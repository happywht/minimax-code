"""Bounded stored worksheet cells, formula text and unverified cached values.

Only the resource-limited child calls this module. Nothing is saved, calculated,
refreshed or dereferenced. These checks are not a filesystem/network sandbox.
"""

from __future__ import annotations

import io
import json
import math
import posixpath
import re
from datetime import date, datetime, time

from .attachment_docx import Reject, package_parts, require, safe_name
from .attachment_metadata import PDF_SCOPE, failure

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"
OFFICE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
MAIN = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
SHEET = "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"
MAX_SHEETS = 32
MAX_CELLS = 100_000
MAX_STRINGS = 50_000
MAX_STYLES = 4096
MAX_ELEMENTS = 1_000_000
MAX_VALUE_BYTES = 65536
MAX_DISPLAY = 1500
ROWS, COLUMNS = 100, 20
COORDINATE = re.compile(r"([A-Z]{1,3})([1-9][0-9]{0,6})\Z")
NUMBER = re.compile(r"[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[Ee][+-]?[0-9]+)?\Z")
ISO_DATE = re.compile(r"[0-9]{4}-[0-9]{2}-[0-9]{2}\Z")
ISO_TIME = re.compile(r"[0-9]{2}:[0-9]{2}(?::[0-9]{2}(?:\.[0-9]{1,3})?)?Z?\Z")
INERT_RELS = {
    OFFICE + value
    for value in (
        "officeDocument",
        "worksheet",
        "styles",
        "sharedStrings",
        "theme",
        "drawing",
        "chart",
        "image",
        "comments",
        "vmlDrawing",
        "table",
        "calcChain",
        "extended-properties",
        "custom-properties",
    )
} | {REL + "/metadata/core-properties"}


def _failure(reason, *, unavailable=False):
    return failure(reason, unavailable=unavailable, scope=PDF_SCOPE)


def _text(value):
    require(isinstance(value, str) and "\x00" not in value)
    require(len(value.encode("utf8")) <= MAX_VALUE_BYTES, "package_limit")
    return value


def _plain(element):
    if element is None:
        return ""
    chunks = []
    for child in element:
        if child.tag == f"{{{NS}}}t":
            require(not list(child))
            chunks.append(child.text or "")
        elif child.tag == f"{{{NS}}}r":
            require(
                all(item.tag in {f"{{{NS}}}rPr", f"{{{NS}}}t"} for item in child),
                "unsupported_feature",
            )
            chunks.extend(item.text or "" for item in child if item.tag == f"{{{NS}}}t")
        else:
            require(False, "unsupported_feature")
    return _text("".join(chunks))


def _coordinate(value):
    match = COORDINATE.fullmatch(value or "")
    require(match is not None)
    column = 0
    for char in match[1]:
        column = column * 26 + ord(char) - 64
    row = int(match[2])
    require(column <= 16384 and row <= 1048576)
    return row, column


def _stored(kind, raw, strings):
    if raw is None:
        return None
    _text(raw)
    if kind == "s":
        require(raw.isascii() and raw.isdecimal() and len(raw) <= 8)
        index = int(raw)
        require(index < len(strings))
        return strings[index]
    if kind == "n":
        require(len(raw) <= 1024 and NUMBER.fullmatch(raw) is not None)
        require(math.isfinite(float(raw)))
    elif kind == "b":
        require(raw in {"0", "1"})
    elif kind == "e":
        require(raw.startswith("#") and len(raw) <= 64)
    elif kind == "d":
        # Validate every stored ISO value, even outside the display window.
        # Keep the original text; no serial-date/style inference is displayed.
        if ISO_DATE.fullmatch(raw):
            date.fromisoformat(raw)
        elif ISO_TIME.fullmatch(raw):
            time.fromisoformat(raw.removesuffix("Z"))
        else:
            pieces = raw.split("T")
            require(len(pieces) == 2 and ISO_DATE.fullmatch(pieces[0]) is not None)
            require(ISO_TIME.fullmatch(pieces[1]) is not None)
            datetime.fromisoformat(raw.removesuffix("Z"))
    elif kind != "str":
        require(False, "unsupported_feature")
    return raw


def _xml(parts, fromstring):
    trees, elements = {}, 0
    # Parse every XML/rels member, including unused parts after display truncation.
    for name, value in parts.items():
        if name.endswith((".xml", ".rels")):
            tree = fromstring(value, forbid_dtd=True, forbid_entities=True, forbid_external=True)
            elements += sum(1 for _ in tree.iter())
            require(elements <= MAX_ELEMENTS, "package_limit")
            trees[name] = tree
    require(
        "[Content_Types].xml" in trees
        and "xl/workbook.xml" in trees
        and "word/document.xml" not in parts
    )
    types = trees["[Content_Types].xml"]
    require(types.tag == f"{{{CT}}}Types")
    defaults, overrides = {}, {}
    for item in types:
        mime = item.get("ContentType", "")
        require(
            mime
            and not any(
                value in mime.lower()
                for value in (
                    "macroenabled",
                    "vbaproject",
                    "oleobject",
                    "activex",
                    "officedocument.package",
                    "externallink",
                    "connections",
                    "chartsheet",
                )
            ),
            "unsupported_feature",
        )
        if item.tag == f"{{{CT}}}Override":
            name = item.get("PartName", "")
            require(name.startswith("/"))
            safe_name(name[1:])
            require(name[1:] in parts and name[1:] not in overrides)
            overrides[name[1:]] = mime
        elif item.tag == f"{{{CT}}}Default":
            extension = item.get("Extension", "")
            require(extension and extension not in defaults)
            require("spreadsheetml" not in mime, "unsupported_feature")
            defaults[extension] = mime
        else:
            require(False)
    require(overrides.get("xl/workbook.xml") == MAIN, "unsupported_feature")
    for name, value in parts.items():
        require(
            not any(
                token in name.casefold()
                for token in (
                    "vbaproject",
                    "activex",
                    "externallinks/",
                    "connections",
                    "querytables/",
                    "embeddings/",
                    "macrosheets/",
                    "dialogsheets/",
                )
            ),
            "unsupported_feature",
        )
        if name.endswith("/") or name == "[Content_Types].xml":
            continue
        mime = overrides.get(name, defaults.get(name.rsplit(".", 1)[-1], ""))
        require(mime)
        if name not in trees and (mime.endswith("+xml") or mime.endswith("/xml")):
            tree = fromstring(value, forbid_dtd=True, forbid_entities=True, forbid_external=True)
            elements += sum(1 for _ in tree.iter())
            require(elements <= MAX_ELEMENTS, "package_limit")
            trees[name] = tree
    relationships = {}
    for name, tree in trees.items():
        if not name.endswith(".rels"):
            continue
        require(tree.tag == f"{{{REL}}}Relationships")
        parent, leaf = posixpath.split(name)
        source = (
            "" if name == "_rels/.rels" else posixpath.join(posixpath.dirname(parent), leaf[:-5])
        )
        require(name == "_rels/.rels" or posixpath.basename(parent) == "_rels" and source in parts)
        ids = {}
        for item in tree:
            require(item.tag == f"{{{REL}}}Relationship")
            identity, kind, target = item.get("Id"), item.get("Type"), item.get("Target", "")
            require(identity and identity not in ids and target)
            require(item.get("TargetMode") in {None, "Internal"}, "unsupported_feature")
            require(kind in INERT_RELS, "unsupported_feature")
            require(
                not any(char in target for char in "\\:%?#") and not target.startswith("//"),
                "unsafe_package",
            )
            resolved = (
                target[1:]
                if target.startswith("/")
                else posixpath.normpath(posixpath.join(posixpath.dirname(source), target))
            )
            safe_name(resolved)
            require(resolved in parts)
            ids[identity] = (kind, resolved)
        relationships[source] = ids
    require(
        [
            item
            for item in relationships.get("", {}).values()
            if item[0] == OFFICE + "officeDocument"
        ]
        == [(OFFICE + "officeDocument", "xl/workbook.xml")]
    )
    require(
        {name for name, mime in overrides.items() if mime.endswith(".sharedStrings+xml")}
        <= {"xl/sharedStrings.xml"},
        "unsupported_feature",
    )
    return trees, overrides, relationships


def _preflight(parts, fromstring, *, full_sheet=None):
    trees, overrides, relationships = _xml(parts, fromstring)
    strings = []
    if "xl/sharedStrings.xml" in trees:
        shared = trees["xl/sharedStrings.xml"]
        require(shared.tag == f"{{{NS}}}sst")
        require(len(shared) <= MAX_STRINGS, "package_limit")
        for item in shared:
            require(item.tag == f"{{{NS}}}si")
            strings.append(_plain(item))
    styles = trees.get("xl/styles.xml")
    require(styles is None or styles.tag == f"{{{NS}}}styleSheet")
    xfs = styles.find(f"{{{NS}}}cellXfs") if styles is not None else None
    style_count = len(xfs) if xfs is not None else 1
    require(style_count <= MAX_STYLES, "package_limit")
    require(xfs is None or all(item.tag == f"{{{NS}}}xf" for item in xfs))
    workbook = trees["xl/workbook.xml"]
    require(workbook.tag == f"{{{NS}}}workbook", "unsupported_feature")
    props = workbook.find(f"{{{NS}}}workbookPr")
    epoch = "1904" if props is not None and props.get("date1904") in {"1", "true"} else "1900"
    sheet_sets = workbook.findall(f"{{{NS}}}sheets")
    require(len(sheet_sets) == 1)
    sheets = sheet_sets[0]
    require(1 <= len(sheets) <= MAX_SHEETS, "package_limit")
    result, targets, names = [], set(), set()
    total, populated, displayed, clipped = 0, 0, 0, False
    for sheet in sheets:
        require(sheet.tag == f"{{{NS}}}sheet")
        name, state = _text(sheet.get("name", "")), sheet.get("state", "visible")
        require(
            name and name.casefold() not in names and state in {"visible", "hidden", "veryHidden"}
        )
        names.add(name.casefold())
        relationship = relationships.get("xl/workbook.xml", {}).get(
            sheet.get(f"{{{OFFICE[:-1]}}}id")
        )
        require(
            relationship is not None and relationship[0] == OFFICE + "worksheet",
            "unsupported_feature",
        )
        target = relationship[1]
        require(target not in targets and overrides.get(target) == SHEET)
        targets.add(target)
        tree = trees.get(target)
        require(tree is not None and tree.tag == f"{{{NS}}}worksheet", "unsupported_feature")
        selected = {}
        capture = full_sheet is not None and name == full_sheet
        full_cells = {} if capture else None
        max_row, max_column = 0, 0
        if capture:
            require(state == "visible", "hidden_sheet")
            for feature in ("mergeCells", "drawing", "legacyDrawing", "tableParts", "hyperlinks",
                            "conditionalFormatting", "dataValidations", "autoFilter", "sheetProtection"):
                require(tree.find(f"{{{NS}}}{feature}") is None, "complex_sheet")
            for cols in tree.findall(f"{{{NS}}}cols"):
                require(all(item.get("hidden") not in {"1", "true"} for item in cols), "hidden_columns")
        data_sets = tree.findall(f"{{{NS}}}sheetData")
        require(len(data_sets) == 1)
        data = data_sets[0]
        require(sum(1 for _ in tree.iter(f"{{{NS}}}c")) == sum(len(row) for row in data))
        previous_row = 0
        for row in data:
            require(
                row.tag == f"{{{NS}}}row"
                and (row.get("r") or "").isascii()
                and (row.get("r") or "").isdecimal()
                and len(row.get("r")) <= 7
            )
            row_number = int(row.get("r"))
            if capture:
                require(row.get("hidden") not in {"1", "true"}, "hidden_rows")
                require(row_number <= 1001, "import_limit")
                max_row = max(max_row, row_number)
            require(previous_row < row_number <= 1048576)
            previous_row = row_number
            previous_column = 0
            for cell in row:
                total += 1
                require(total <= MAX_CELLS, "package_limit")
                require(cell.tag == f"{{{NS}}}c")
                coordinate = cell.get("r", "")
                actual_row, column = _coordinate(coordinate)
                if capture:
                    require(column <= 20, "import_limit")
                    max_column = max(max_column, column)
                require(actual_row == row_number and column > previous_column)
                previous_column = column
                style = cell.get("s", "0")
                require(
                    style.isascii()
                    and style.isdecimal()
                    and len(style) <= 8
                    and int(style) < style_count
                )
                kind = cell.get("t", "n")
                require(
                    kind in {"n", "s", "inlineStr", "b", "e", "str", "d"}, "unsupported_feature"
                )
                require(
                    all(item.tag in {f"{{{NS}}}v", f"{{{NS}}}f", f"{{{NS}}}is"} for item in cell),
                    "unsupported_feature",
                )
                formulae, values, inline = (
                    cell.findall(f"{{{NS}}}{key}") for key in ("f", "v", "is")
                )
                require(len(formulae) <= 1 and len(values) <= 1 and len(inline) <= 1)
                value = values[0] if values else None
                require(value is None or not list(value))
                raw = value.text or "" if value is not None else None
                if capture:
                    require(not formulae, "formula")
                    require(kind not in {"e", "d"}, "unsupported_cell")
                    if kind == "n" and value is not None and raw not in {None, ""} and xfs is not None:
                        require(xfs[int(style)].get("numFmtId", "0") == "0", "formatted_number")
                if formulae:
                    formula = formulae[0]
                    require(
                        formula.get("t") in {None, "normal"} and not list(formula),
                        "unsupported_feature",
                    )
                    text = _text(formula.text or "")
                    require(bool(text.strip()))
                    require(kind not in {"s", "inlineStr"} and not inline, "unsupported_feature")
                    if raw is not None and raw != "":
                        _stored(kind, raw, strings)
                    elif raw == "":
                        require(kind in {"str", "n"})  # Empty numeric <v/> is an unknown cache.
                    record = {
                        "cell": coordinate,
                        "kind": "formula",
                        "stored_type": kind,
                        "formula": text,
                        "cached_present": value is not None,
                        "cached_raw": raw,
                        "cached_state": "missing"
                        if value is None
                        else "unknown"
                        if raw == "" and kind == "n"
                        else "stored",
                        "cache_verified": False,
                        "recalculated": False,
                    }
                    has_value = True
                else:
                    require(not inline or kind == "inlineStr")
                    if kind == "inlineStr":
                        require(value is None)
                        decoded = _plain(inline[0] if inline else None)
                    else:
                        require(not inline)
                        decoded = _stored(kind, raw, strings)
                    has_value = decoded is not None and decoded != ""
                    record = {
                        "cell": coordinate,
                        "kind": "string"
                        if kind in {"s", "inlineStr", "str"}
                        else {"n": "number", "b": "boolean", "e": "error", "d": "iso_date"}[kind],
                        "stored_type": kind,
                        "raw": decoded,
                    }
                if capture and has_value:
                    require(record["kind"] in {"string", "number", "boolean"}, "unsupported_cell")
                    require(len(record["raw"]) <= 4096, "import_limit")
                    full_cells[(actual_row, column)] = record
                if has_value:
                    populated += 1
                    if actual_row <= ROWS and column <= COLUMNS and displayed < MAX_DISPLAY:
                        selected[coordinate] = record
                        displayed += 1
                    else:
                        clipped = True
        item = {"name": name, "state": state, "epoch": epoch, "selected": selected}
        if capture:
            require(max_row >= 1 and max_column >= 1 and max_row * max_column <= 5000, "import_limit")
            item["full_cells"] = full_cells
            item["full_rows"] = max_row
            item["full_columns"] = max_column
        result.append(item)
    require(targets == {name for name, mime in overrides.items() if mime == SHEET})
    return result, populated, clipped



def extract_sheet(data, sheet_name):
    """Return complete bounded raw cells, never computed or cached formulas."""
    try:
        from defusedxml.ElementTree import fromstring

        require(isinstance(sheet_name, str) and 1 <= len(sheet_name) <= 31, "sheet_name")
        parts = package_parts(data)
        sheets, _, _ = _preflight(parts, fromstring, full_sheet=sheet_name)
        selected = next((sheet for sheet in sheets if sheet["name"] == sheet_name), None)
        require(selected is not None, "sheet_not_found")
        cells = selected["full_cells"]
        rows = []
        for row in range(1, selected["full_rows"] + 1):
            values = []
            for column in range(1, selected["full_columns"] + 1):
                record = cells.get((row, column))
                if record is None:
                    values.append("")
                elif record["kind"] == "boolean":
                    values.append("true" if record["raw"] == "1" else "false")
                else:
                    values.append(record["raw"])
            rows.append(values)
        result = {"ok": True, "rows": rows}
        require(len(json.dumps(result, ensure_ascii=False).encode("utf-8")) <= 1024 * 1024,
                "import_limit")
        return result
    except ImportError:
        return {"ok": False, "reason": "missing_dependency"}
    except Reject as error:
        return {"ok": False, "reason": error.reason}
    except (MemoryError, RecursionError):
        return {"ok": False, "reason": "import_limit"}
    except Exception:
        return {"ok": False, "reason": "invalid_package"}

def parse_xlsx(data):
    try:
        from defusedxml.ElementTree import fromstring
        from openpyxl import load_workbook
    except ImportError:
        return _failure("missing_dependency", unavailable=True)
    try:
        parts = package_parts(data)
        sheets, populated, clipped = _preflight(parts, fromstring)
        # No workbook writer, formula engine or external-link APIs. Both readers
        # use these same bytes; raw preflight values remain authoritative.
        books = []
        output, size, lines = [], 0, 0

        def emit(value):
            nonlocal size, lines, clipped
            line = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
            line = (
                line.replace("\u0085", "\\u0085")
                .replace("\u2028", "\\u2028")
                .replace("\u2029", "\\u2029")
                + "\n"
            )
            if size + len(line.encode("utf8")) > 65536 or lines >= 2000:
                clipped = True
                return
            output.append(line)
            size += len(line.encode("utf8"))
            lines += 1

        try:
            for data_only in (False, True):
                books.append(
                    load_workbook(
                        io.BytesIO(data),
                        read_only=True,
                        data_only=data_only,
                        keep_vba=False,
                        keep_links=False,
                        rich_text=False,
                    )
                )
            require(all(book.sheetnames == [sheet["name"] for sheet in sheets] for book in books))
            for index, sheet in enumerate(sheets, 1):
                emit(
                    {
                        "sheet": index,
                        "name": sheet["name"],
                        "state": sheet["state"],
                        "epoch": sheet["epoch"],
                        "window": "A1:T100",
                        "values": "stored_raw",
                    }
                )
                first, cached = (book[sheet["name"]] for book in books)
                require(first.sheet_state == cached.sheet_state == sheet["state"])
                first.reset_dimensions()
                cached.reset_dimensions()
                seen = set()
                formula_rows = first.iter_rows(min_row=1, max_row=ROWS, min_col=1, max_col=COLUMNS)
                cache_rows = cached.iter_rows(min_row=1, max_row=ROWS, min_col=1, max_col=COLUMNS)
                for row_number in range(1, ROWS + 1):
                    left, right = next(formula_rows, None), next(cache_rows, None)
                    require((left is None) == (right is None))
                    if left is None:
                        break
                    require(len(left) == len(right) == COLUMNS)
                    for column, (cell, cache) in enumerate(zip(left, right, strict=True), 1):
                        coordinate = chr(64 + column) + str(row_number)
                        record = sheet["selected"].get(coordinate)
                        if record is None:
                            continue
                        require(cell.coordinate == coordinate and cache.coordinate == coordinate)
                        require((cell.data_type == "f") == (record["kind"] == "formula"))
                        seen.add(coordinate)
                        emit(record)
                require(seen == set(sheet["selected"]))
        finally:
            for book in books:
                book.close()
        status = "empty" if not populated else "truncated" if clipped else "parsed"
        return {
            "parser_status": status,
            "parser_reason": "preview_limit" if status == "truncated" else "ok",
            "parser_scope": PDF_SCOPE,
            "parser_version": 2,
            "preview_text": "".join(output) if populated else "",
            "preview_truncated": int(status == "truncated"),
            "quality_status": "unverified",
        }
    except Reject as error:
        return _failure(error.reason)
    except (MemoryError, RecursionError):
        return _failure("package_limit")
    except Exception:
        # No exception text, filenames, cell snippets, paths or unbounded fallback.
        return _failure("invalid_package")
