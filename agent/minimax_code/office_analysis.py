"""Static aggregation and a separate chart-bearing XLSX verification scope."""
from __future__ import annotations

import hashlib
import io
from decimal import ROUND_HALF_UP, Decimal
from typing import Literal

from pydantic import BaseModel, ConfigDict, Field, model_validator

from .office_xlsx import XlsxSpec, canonical, generate, text


class Summary(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    column: str
    operation: Literal["sum", "count", "average"]
    decimals: int = Field(default=2, ge=0, le=8)


class Chart(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    kind: Literal["bar", "line"]
    category_column: str
    value_column: str
    title: str = Field(min_length=1, max_length=100)


class AnalysisSpec(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    spec: XlsxSpec
    chart: Chart
    source_rows: int = Field(ge=1, le=1000)

    @model_validator(mode="after")
    def check_chart(self):
        if self.source_rows > len(self.spec.rows):
            raise ValueError("Chart source rows exceed table")
        names = {column.name: column for column in self.spec.columns}
        category, value = names.get(self.chart.category_column), names.get(self.chart.value_column)
        if category is None or value is None or category.kind not in {"text", "number"} or value.kind != "number":
            raise ValueError("Chart requires an existing text/number category and numeric value column")
        text(self.chart.title)
        if any(ord(char) < 32 for char in self.chart.title):
            raise ValueError("Chart title cannot contain control characters")
        return self


def materialize(spec: XlsxSpec, summaries: list[Summary]) -> XlsxSpec:
    if not summaries:
        return spec
    if len(summaries) > 20 or len(spec.columns) > 18:
        raise ValueError("Summaries support at most 20 operations and 18 input columns")
    names = {column.name: index for index, column in enumerate(spec.columns)}
    if any(name.casefold() in {column.name.casefold() for column in spec.columns}
           for name in ("汇总项目", "汇总值")):
        raise ValueError("Input columns cannot be named 汇总项目 or 汇总值 when summaries are requested")
    rows = [row + [None, None] for row in spec.rows]
    for summary in summaries:
        if summary.column not in names:
            raise ValueError("Summary column does not exist")
        index = names[summary.column]
        column = spec.columns[index]
        values = [row[index] for row in spec.rows if row[index] is not None]
        if summary.operation != "count" and column.kind != "number":
            raise ValueError("sum and average require a numeric column")
        if summary.operation == "count":
            number = len(values)
        elif summary.operation == "sum":
            result = sum((Decimal(str(value)) for value in values), Decimal(0))
            number = int(result) if result == result.to_integral_value() else float(result)
        elif not values:
            number = None
        else:
            result = (sum((Decimal(str(value)) for value in values), Decimal(0)) / len(values)).quantize(
                Decimal(1).scaleb(-summary.decimals), rounding=ROUND_HALF_UP)
            number = int(result) if result == result.to_integral_value() else float(result)
        if number is not None and summary.operation != "count":
            if Decimal(str(number)) != result:
                raise ValueError("Summary cannot be represented safely in a spreadsheet number")
        label = f"{summary.operation}({summary.column})"
        if summary.operation == "average":
            label += f" [小数位 {summary.decimals}, 四舍五入]"
        rows.append([None] * len(spec.columns) + [label, number])
    return XlsxSpec.model_validate({"sheet_name": spec.sheet_name, "columns": [
        *[column.model_dump() for column in spec.columns],
        {"name": "汇总项目", "kind": "text"}, {"name": "汇总值", "kind": "number"}], "rows": rows})


def generate_chart(analysis: AnalysisSpec) -> bytes:
    from openpyxl import load_workbook
    from openpyxl.chart import BarChart, LineChart, Reference

    data = generate(analysis.spec)
    book = load_workbook(io.BytesIO(data))
    try:
        sheet = book.active
        names = [column.name for column in analysis.spec.columns]
        category = names.index(analysis.chart.category_column) + 1
        value = names.index(analysis.chart.value_column) + 1
        chart = BarChart() if analysis.chart.kind == "bar" else LineChart()
        chart.title = analysis.chart.title
        chart.add_data(Reference(sheet, min_col=value, min_row=1, max_row=analysis.source_rows + 1),
                       titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=category, min_row=2, max_row=analysis.source_rows + 1))
        sheet.add_chart(chart, "A" + str(len(analysis.spec.rows) + 4))
        output = io.BytesIO()
        book.save(output)
        data = output.getvalue()
        if not 0 < len(data) <= 2 * 1024 * 1024:
            raise ValueError("Chart XLSX exceeds 2 MiB")
        return data
    finally:
        book.close()


def verify_chart(data: bytes, analysis: AnalysisSpec) -> dict:
    """Separate generated chart scope: exact package, cells and series ranges."""
    from defusedxml.ElementTree import fromstring
    from openpyxl import load_workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.utils import get_column_letter, quote_sheetname

    from .attachment_docx import package_parts, require
    from .attachment_xlsx import _preflight
    from .office_xlsx import PARTS

    require(0 < len(data) <= 2 * 1024 * 1024)
    parts = package_parts(data)
    require(set(parts) == PARTS | {"xl/charts/chart1.xml", "xl/drawings/drawing1.xml",
                                   "xl/drawings/_rels/drawing1.xml.rels",
                                   "xl/worksheets/_rels/sheet1.xml.rels"})
    require(sum(map(len, parts.values())) <= 8 * 1024 * 1024)
    sheets, _, _ = _preflight(parts, fromstring)
    require(len(sheets) == 1 and sheets[0]["name"] == analysis.spec.sheet_name)
    book = load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
    try:
        require(book.sheetnames == [analysis.spec.sheet_name])
        sheet = book.active
        matrix = [[column.name for column in analysis.spec.columns], *analysis.spec.rows]
        require(sheet.max_row == len(matrix) and sheet.max_column == len(analysis.spec.columns))
        for row in sheet.iter_rows():
            for cell in row:
                expected = matrix[cell.row - 1][cell.column - 1]
                require(cell.data_type != "f")
                if type(expected) in {int, float}:
                    require(type(cell.value) in {int, float} and Decimal(str(cell.value)) == Decimal(str(expected)))
                else:
                    require(type(cell.value) is type(expected) and cell.value == expected)
        require(len(sheet._charts) == 1)
        chart = sheet._charts[0]
        require(isinstance(chart, BarChart if analysis.chart.kind == "bar" else LineChart))
        require(len(chart.series) == 1)
        names = [column.name for column in analysis.spec.columns]
        category, value = names.index(analysis.chart.category_column) + 1, names.index(analysis.chart.value_column) + 1
        expected_c = str(Reference(sheet, min_col=category, min_row=2, max_row=analysis.source_rows + 1))
        expected_v = str(Reference(sheet, min_col=value, min_row=2, max_row=analysis.source_rows + 1))
        require(chart.series[0].cat is not None and chart.series[0].cat.numRef.f == expected_c)
        require(chart.series[0].val is not None and chart.series[0].val.numRef.f == expected_v)
        require(chart.series[0].tx is not None and chart.series[0].tx.strRef is not None
                and chart.series[0].tx.strRef.f == f"{quote_sheetname(sheet.title)}!{get_column_letter(value)}1")
        require(chart.title.tx.rich.p[0].r[0].t == analysis.chart.title)
    finally:
        book.close()
    return {"scope": "xlsx_chart_values_v1", "input_sha256": hashlib.sha256(canonical(analysis.model_dump())).hexdigest(),
            "file_sha256": hashlib.sha256(data).hexdigest(), "size_bytes": len(data),
            "cell_check": "passed", "chart_reference_check": "passed", "render_status": "not_attempted"}
