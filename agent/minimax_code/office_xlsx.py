"""Fixed plain-cell XLSX writer and a separate, complete stored-value verifier.

No templates, formulas, external resources or Office execution. Import the writer
only in the resource-limited worker; CLI verification never accesses application DBs.
"""
from __future__ import annotations

import hashlib
import io
import json
import math
import re
from decimal import Decimal
from typing import Any, Literal

from pydantic import BaseModel, ConfigDict, Field, model_validator

MAX_INPUT = 1024 * 1024
MAX_OUTPUT = 2 * 1024 * 1024
SCOPE = "xlsx_plain_cells_v1"
WRITER_VERSION = "3.1.5"
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
PARTS = {
    "docProps/app.xml", "docProps/core.xml", "xl/theme/theme1.xml",
    "xl/worksheets/sheet1.xml", "xl/styles.xml", "_rels/.rels", "xl/workbook.xml",
    "xl/_rels/workbook.xml.rels", "[Content_Types].xml",
}


def canonical(value: Any) -> bytes:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"),
                      allow_nan=False).encode("utf-8")


def digest(value: Any) -> str:
    return hashlib.sha256(canonical(value)).hexdigest()


def text(value: str) -> str:
    if any(not (char in "\t\n\r" or 32 <= ord(char) <= 0xD7FF
                or 0xE000 <= ord(char) <= 0xFFFD or 0x10000 <= ord(char) <= 0x10FFFF)
           for char in value):
        raise ValueError("Office text contains unsupported XML characters")
    return value


def filename(value: str) -> str:
    text(value)
    if (not 6 <= len(value) <= 120 or not value.endswith(".xlsx")
            or any(char in value for char in '\\/:*?"<>|')
            or any(ord(char) < 32 or ord(char) == 127 for char in value)
            or value != value.strip() or value.endswith(" .xlsx")
            or value.split(".")[0].upper() in {"CON", "PRN", "AUX", "NUL",
                *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}):
        raise ValueError("Use a plain XLSX display filename")
    return value


class Column(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    name: str = Field(min_length=1, max_length=80)
    kind: Literal["text", "number", "boolean"]

    @model_validator(mode="after")
    def check_name(self):
        text(self.name)
        if not self.name.strip():
            raise ValueError("Column name must contain text")
        return self


class XlsxSpec(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    sheet_name: str = Field(min_length=1, max_length=31)
    columns: list[Column] = Field(min_length=1, max_length=20)
    rows: list[list[Any]] = Field(max_length=1000)

    @model_validator(mode="after")
    def check_cells(self):
        text(self.sheet_name)
        if (any(char in self.sheet_name for char in "[]:*?/\\")
                or any(ord(char) < 32 for char in self.sheet_name)
                or self.sheet_name.startswith("'") or self.sheet_name.endswith("'")
                or not self.sheet_name.strip()):
            raise ValueError("Invalid sheet name")
        if len({column.name.casefold() for column in self.columns}) != len(self.columns):
            raise ValueError("Column names must be distinct")
        if (len(self.rows) + 1) * len(self.columns) > 5000:
            raise ValueError("Office table exceeds 5000 cells including headers")
        normalized = []
        for row in self.rows:
            if len(row) != len(self.columns):
                raise ValueError("Every row must match the column count")
            values = []
            for column, value in zip(self.columns, row, strict=True):
                if value is None:
                    pass
                elif column.kind == "text":
                    if type(value) is not str or len(value) > 4096:
                        raise ValueError("Text cells must be strings up to 4096 characters")
                    text(value)
                    value = value or None  # Explicit API rule: empty text is a blank cell.
                elif column.kind == "boolean":
                    if type(value) is not bool:
                        raise ValueError("Boolean cells require JSON booleans")
                else:
                    if (type(value) not in {int, float} or abs(value) > 999999999999999
                            or not math.isfinite(value)
                            or value == 0 and math.copysign(1, value) < 0):
                        raise ValueError("Number cells require finite JSON numbers")
                    number = Decimal(str(value)).normalize()
                    if (len(number.as_tuple().digits) > 15 or abs(number) > 999999999999999
                            or number and abs(number) < Decimal("1e-12")):
                        raise ValueError("Use text for numbers beyond 15 digits or supported magnitude")
                values.append(value)
            normalized.append(values)
        self.rows = normalized
        if len(canonical(self.model_dump())) > MAX_INPUT:
            raise ValueError("Office structured input exceeds 1 MiB")
        return self


def generate(spec: XlsxSpec) -> bytes:
    import importlib.metadata

    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    if importlib.metadata.version("openpyxl") != WRITER_VERSION:
        raise ValueError("Unsupported Office writer version")
    book = Workbook()
    try:
        book.properties.creator = "MiniMax Code"
        book.properties.lastModifiedBy = "MiniMax Code"
        sheet = book.active
        sheet.title = spec.sheet_name
        sheet.freeze_panes = "A2"
        for index in range(1, len(spec.columns) + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 20
        matrix = [[column.name for column in spec.columns], *spec.rows]
        for row, values in enumerate(matrix, 1):
            for column, value in enumerate(values, 1):
                if value is None:
                    continue
                cell = sheet.cell(row, column, value)
                if type(value) is str:
                    cell.data_type = "s"  # Never allow '=...' or '#N/A' inference.
                    cell.number_format = "@"
                if row == 1:
                    cell.font = Font(bold=True)
        output = io.BytesIO()
        book.save(output)
        data = output.getvalue()
        if not 0 < len(data) <= MAX_OUTPUT:
            raise ValueError("Office output exceeds 2 MiB")
        return data
    finally:
        book.close()


def verify(data: bytes, spec: XlsxSpec) -> dict:
    """Read actual ZIP/XML and compare every cell, independently of writer state."""
    from defusedxml.ElementTree import fromstring

    from .attachment_docx import package_parts, require
    from .attachment_xlsx import _preflight

    require(0 < len(data) <= MAX_OUTPUT, "package_limit")
    parts = package_parts(data)
    require(set(parts) == PARTS, "unsupported_feature")
    require(sum(map(len, parts.values())) <= 8 * 1024 * 1024
            and max(map(len, parts.values())) <= 4 * 1024 * 1024, "package_limit")
    elements = 0
    for raw in parts.values():
        root = fromstring(raw, forbid_dtd=True, forbid_entities=True, forbid_external=True)
        for node in root.iter():
            elements += 1
            require(elements <= 100000, "package_limit")
            require(node.tag != f"{{{NS}}}f", "unsupported_feature")
    sheets, _, _ = _preflight(parts, fromstring)
    require(len(sheets) == 1 and sheets[0]["name"] == spec.sheet_name
            and sheets[0]["state"] == "visible")
    workbook = fromstring(parts["xl/workbook.xml"])
    names = workbook.find(f"{{{NS}}}definedNames")
    require(names is None or len(names) == 0, "unsupported_feature")
    worksheet = fromstring(parts["xl/worksheets/sheet1.xml"])
    allowed = {"worksheet", "sheetPr", "outlinePr", "pageSetUpPr", "dimension", "sheetViews",
               "sheetView", "pane", "selection", "sheetFormatPr", "cols", "col", "sheetData",
               "row", "c", "pageMargins", "is", "t", "v"}
    require(all(node.tag in {f"{{{NS}}}{tag}" for tag in allowed}
                for node in worksheet.iter()), "unsupported_feature")
    matrix = [[column.name for column in spec.columns], *spec.rows]
    expected = {(row, column): value for row, values in enumerate(matrix, 1)
                for column, value in enumerate(values, 1) if value is not None}
    seen = set()
    for cell in worksheet.iter(f"{{{NS}}}c"):
        coordinate = cell.get("r", "")
        match = re.fullmatch(r"([A-T])([1-9][0-9]{0,3})", coordinate)
        require(match is not None)
        key = int(match[2]), ord(match[1]) - 64
        require(key in expected and key not in seen)
        seen.add(key)
        value = expected[key]
        if type(value) is str:
            require(cell.get("t") == "inlineStr" and len(cell) == 1)
            inline = cell.find(f"{{{NS}}}is")
            require(inline is not None and len(inline) == 1
                    and inline[0].tag == f"{{{NS}}}t" and len(inline[0]) == 0)
            require((inline[0].text or "") == value)
        else:
            require(len(cell) == 1 and cell[0].tag == f"{{{NS}}}v")
            raw = cell[0].text
            if type(value) is bool:
                require(cell.get("t") == "b" and raw == str(int(value)))
            else:
                require(cell.get("t", "n") == "n" and raw is not None)
                require(Decimal(raw).is_finite() and Decimal(raw) == Decimal(str(value)))
    require(seen == set(expected))
    return {
        "version": 1, "scope": SCOPE, "input_sha256": digest(spec.model_dump()),
        "file_sha256": hashlib.sha256(data).hexdigest(), "size_bytes": len(data),
        "sheet_name": spec.sheet_name, "data_rows": len(spec.rows),
        "column_count": len(spec.columns), "cell_count": len(matrix) * len(spec.columns),
        "formula_count": 0, "external_relationship_count": 0, "macro_count": 0,
        "structure_status": "passed", "business_quality": "unverified",
        "render_status": "not_attempted", "writer": {"name": "openpyxl", "version": WRITER_VERSION},
        "validator_version": 1,
    }


def main():
    import argparse
    from pathlib import Path

    from .text_artifacts import _read

    parser = argparse.ArgumentParser(description="Read-only full plain-cell XLSX verification")
    parser.add_argument("operation", choices=["verify"])
    parser.add_argument("file")
    parser.add_argument("--spec", required=True, help="Original normalized table specification JSON")
    args = parser.parse_args()
    try:
        raw = _read(Path(args.spec).absolute())
        if len(raw) > MAX_INPUT:
            raise ValueError("input_limit")
        spec = XlsxSpec.model_validate_json(raw)
        result = verify(_read(Path(args.file).absolute()), spec)
    except Exception:
        print(json.dumps({"error": "XLSX verification failed"}))
        return 1
    print(json.dumps(result, ensure_ascii=True, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
