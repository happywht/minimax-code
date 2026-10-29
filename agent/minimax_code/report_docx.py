"""Generate bounded DOCX reports from structured operator-supplied text."""
from __future__ import annotations

import threading
from io import BytesIO

MAX_OUTPUT_BYTES = 2 * 1024 * 1024
_GATE = threading.BoundedSemaphore(1)


def generate_document(filename: str, title: str, paragraphs=None, sections=None, sources=None) -> bytes:
    if not _GATE.acquire(blocking=False):
        raise ValueError("DOCX generation is busy; retry later")
    try:
        return _generate_document(filename, title, paragraphs, sections, sources)
    finally:
        _GATE.release()


def _generate_document(filename: str, title: str, paragraphs, sections, sources) -> bytes:
    from .report_content import report_blocks

    blocks = report_blocks(filename, title, paragraphs, sections, ".docx", sources)
    try:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise ValueError("DOCX generation requires the attachments dependency") from exc
    document = Document()
    for kind, value in blocks:
        if kind == "heading0":
            document.add_heading(value, level=0)
        elif kind == "heading1":
            document.add_heading(value, level=1)
        elif kind == "paragraph":
            document.add_paragraph(value)
        else:
            table = document.add_table(rows=len(value), cols=len(value[0]))
            table.style = "Table Grid"
            for document_row, row in zip(table.rows, value, strict=True):
                for document_cell, cell in zip(document_row.cells, row, strict=True):
                    document_cell.text = cell
    output = BytesIO()
    document.save(output)
    data = output.getvalue()
    if not 0 < len(data) <= MAX_OUTPUT_BYTES:
        raise ValueError("Generated DOCX exceeds 2 MiB")
    # Inspect the actual saved document in body order, including every cell.
    try:
        reopened = Document(BytesIO(data))
        actual = []
        for element in reopened.element.body:
            if element.tag == qn("w:p"):
                paragraph = Paragraph(element, reopened)
                style = paragraph.style.name
                kind = "heading0" if style == "Title" else "heading1" if style == "Heading 1" else "paragraph"
                actual.append((kind, paragraph.text))
            elif element.tag == qn("w:tbl"):
                table = Table(element, reopened)
                actual.append(("table", [[cell.text for cell in row.cells] for row in table.rows]))
            elif element.tag != qn("w:sectPr"):
                raise ValueError("Unexpected document body element")
    except Exception as exc:
        raise ValueError("Generated DOCX could not be verified") from exc
    if actual != blocks:
        raise ValueError("Generated DOCX body did not round-trip")
    return data
