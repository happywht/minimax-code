"""Bounded PDF/DOCX original reads beyond stored previews."""
from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from reportlab.pdfgen import canvas

from minimax_code.agent.tools import ToolRegistry
from minimax_code.agent.tools.project_attachments import register_project_attachment_tools
from minimax_code.attachments import AttachmentService
from minimax_code.execution import ExecutionPolicy, execution_scope
from minimax_code.storage.dao.projects import ProjectsDAO
from minimax_code.storage.dao.sessions import SessionsDAO
from minimax_code.storage.db import AsyncDatabase


@pytest.fixture
async def env(tmp_path):
    db = AsyncDatabase(tmp_path / "data" / "test.db")
    await db.connect()
    await db.migrate()
    await ProjectsDAO(db).create(id="p", name="P", root_path=str(tmp_path))
    await ProjectsDAO(db).create(id="other", name="Other", root_path=str(tmp_path))
    sessions = SessionsDAO(db)
    await sessions.create(id="s", project_id="p")
    service = AttachmentService(db, db.path.parent)
    identity = await service.dao.service_id()

    async def upload(project, filename, data):
        async def chunks():
            yield data
        return await service.upload(project, filename, identity, chunks())

    registry = ToolRegistry()
    register_project_attachment_tools(registry, db, "s", "p")
    yield db, sessions, service, upload, registry
    await db.close()


def make_pdf():
    output = BytesIO()
    pdf = canvas.Canvas(output)
    pdf.drawString(20, 500, "A" * 70000)
    pdf.showPage()
    pdf.drawString(20, 500, "PAGE TWO AFTER PREVIEW")
    pdf.showPage()
    pdf.showPage()
    pdf.save()
    return output.getvalue()


def make_docx():
    document = Document()
    document.add_paragraph("X" * 70000)
    document.add_paragraph("AFTER PREVIEW")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "TABLE CELL"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


async def test_pdf_reads_pages_beyond_preview_and_explicit_empty(env):
    _, _, _, upload, registry = env
    row = await upload("p", "source.pdf", make_pdf())
    assert row["preview_truncated"]
    with execution_scope(ExecutionPolicy(read_only=True)):
        first = await registry.dispatch("read_project_document", {"attachment_id": row["id"]})
    assert first.success and len(first.output["text"]) == 4000
    assert first.output["truncated"] and first.output["next_cursor"] == {"unit": 1, "offset": 4000}
    second = await registry.dispatch("read_project_document", {"attachment_id": row["id"], "unit": 2})
    assert second.success and "PAGE TWO AFTER PREVIEW" in second.output["text"]
    assert second.output["next_cursor"] == {"unit": 3, "offset": 0}
    blank = await registry.dispatch("read_project_document", {"attachment_id": row["id"], "unit": 3})
    assert blank.success and blank.output["status"] == "empty" and blank.output["ocr"] == "not_attempted"
    assert "untrusted" in blank.output["untrusted_content_notice"]


async def test_docx_segments_and_session_rebind(env, monkeypatch):
    from minimax_code import attachment_parser

    _, sessions, service, upload, registry = env
    raw = make_docx()
    row = await upload("p", "source.docx", raw)
    assert row["preview_truncated"]
    first = await registry.dispatch("read_project_document", {"attachment_id": row["id"]})
    assert first.success and first.output["next_cursor"] == {"unit": 1, "offset": 4000}
    tail = await registry.dispatch("read_project_document", {"attachment_id": row["id"], "unit": 1, "offset": 68000})
    assert tail.success and len(tail.output["text"]) == 2000 and tail.output["next_cursor"] == {"unit": 2, "offset": 0}
    table = await registry.dispatch("read_project_document", {"attachment_id": row["id"], "unit": 3})
    assert table.success and table.output["text"] == "TABLE CELL" and "/cell1/" in table.output["location"]
    assert (await service.content("p", row["id"]))[1] == raw
    foreign = await upload("other", "foreign.docx", raw)
    assert not (await registry.dispatch("read_project_document", {"attachment_id": foreign["id"]})).success
    assert not (await registry.dispatch("read_project_document", {"attachment_id": row["id"], "project_id": "other"})).success
    original = attachment_parser.read_document_segment

    async def rebind(*args):
        result = await original(*args)
        await sessions.update("s", project_id="other")
        return result

    monkeypatch.setattr(attachment_parser, "read_document_segment", rebind)
    assert not (await registry.dispatch("read_project_document", {"attachment_id": row["id"], "unit": 2})).success


async def test_corrupted_snapshot_is_not_read(env):
    _, _, service, upload, registry = env
    row = await upload("p", "source.docx", make_docx())
    (service.vault / row["snapshot_id"]).write_bytes(b"corrupted")
    result = await registry.dispatch("read_project_document", {"attachment_id": row["id"]})
    assert not result.success and "Attachment unavailable" in result.error
