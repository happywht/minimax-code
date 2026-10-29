"""Full source worksheet extraction, conversion, and access boundaries."""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
from unittest.mock import AsyncMock

import openpyxl
import pytest

from minimax_code.agent.core import AgentCore
from minimax_code.agent.tools import ToolRegistry
from minimax_code.agent.tools.base import ToolResult
from minimax_code.agent.tools.generate_xlsx_from_xlsx import register_xlsx_import_tool
from minimax_code.attachment_parser import extract_xlsx
from minimax_code.attachments import AttachmentService
from minimax_code.execution import ExecutionPolicy, execution_scope
from minimax_code.office_deliveries import OfficeService
from minimax_code.storage.dao.projects import ProjectsDAO
from minimax_code.storage.dao.sessions import SessionsDAO
from minimax_code.storage.db import AsyncDatabase


def workbook_bytes(kind="normal"):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "资料"
    sheet.append(["名称", "数量"])
    for index in range(120):
        sheet.append([f"项目{index}", index])
    if kind == "formula":
        sheet["B101"] = "=1+1"
    elif kind == "date":
        sheet["B3"] = datetime(2024, 1, 1)
    elif kind == "merged":
        sheet.merge_cells("A3:A4")
    elif kind == "hidden":
        sheet.row_dimensions[3].hidden = True
    elif kind == "formatted":
        sheet["B3"].number_format = "0.00%"
    book.create_sheet("另一个")
    output = BytesIO()
    book.save(output)
    return output.getvalue()


@pytest.fixture
async def env(tmp_path):
    db = AsyncDatabase(tmp_path / "data" / "test.db")
    await db.connect()
    await db.migrate()
    projects = ProjectsDAO(db)
    await projects.create(id="p", name="P", root_path=str(tmp_path))
    await projects.create(id="other", name="Other", root_path=str(tmp_path))
    sessions = SessionsDAO(db)
    await sessions.create(id="s", project_id="p")
    service = AttachmentService(db, db.path.parent)
    identity = await service.dao.service_id()

    async def upload(project, data):
        async def chunks():
            yield data
        return await service.upload(project, "input.xlsx", identity, chunks())

    registry = ToolRegistry()
    register_xlsx_import_tool(registry, db, "s", "p")
    yield db, sessions, service, upload, registry
    await db.close()


def request(attachment):
    return {"attachment_id": attachment, "filename": "result.xlsx", "request_key": "key1",
            "source_sheet": "资料", "sheet_name": "结果", "columns": [
                {"source": "数量", "name": "金额", "kind": "number"},
                {"source": "名称", "name": "项目", "kind": "text"}]}


async def test_imports_full_sheet_and_keeps_original_and_bound_project(env):
    db, sessions, service, upload, registry = env
    data = workbook_bytes()
    source = await upload("p", data)
    assert source["preview_truncated"]
    args = request(source["id"])
    args["summaries"] = [{"column": "金额", "operation": "sum"}]
    with execution_scope(ExecutionPolicy(read_only=True)):
        assert not (await registry.dispatch("generate_xlsx_from_xlsx", args)).success
    result = await registry.dispatch("generate_xlsx_from_xlsx", args)
    assert result.success, result.error
    assert result.output["source_sha256"] == source["sha256"]
    office = OfficeService(db, db.path.parent)
    _, output = await office.originals.content("p", result.output["attachment_id"])
    sheet = openpyxl.load_workbook(BytesIO(output)).active
    assert sheet["A1"].value == "金额" and sheet["B121"].value == "项目119"
    assert sheet["D122"].value == sum(range(120))
    assert (await service.content("p", source["id"]))[1] == data
    repeat = await registry.dispatch("generate_xlsx_from_xlsx", args)
    assert repeat.success and repeat.output["delivery_id"] == result.output["delivery_id"]
    await sessions.update("s", project_id="other")
    assert not (await registry.dispatch("generate_xlsx_from_xlsx", args)).success


async def test_import_chart_uses_selected_columns(env):
    _, _, service, upload, registry = env
    source = await upload("p", workbook_bytes())
    args = request(source["id"])
    args["chart"] = {"kind": "bar", "category_column": "项目", "value_column": "金额", "title": "金额图"}
    result = await registry.dispatch("generate_xlsx_from_xlsx", args)
    assert result.success, result.error
    assert result.output["validation_scope"] == "xlsx_chart_values_v1"
    assert "not idempotent" in result.output["notice"]
    _, output = await service.content("p", result.output["attachment_id"])
    assert len(openpyxl.load_workbook(BytesIO(output)).active._charts) == 1


@pytest.mark.parametrize(("kind", "reason"), [("formula", "formula"), ("date", "formatted_number"),
    ("merged", "complex_sheet"), ("hidden", "hidden_rows"), ("formatted", "formatted_number")])
async def test_unsupported_sheet_features_fail_closed(kind, reason):
    result = await extract_xlsx(workbook_bytes(kind), "资料")
    assert result == {"ok": False, "reason": reason}


async def test_missing_and_foreign_attachments_do_not_create_output(env):
    db, _, _, upload, registry = env
    missing = await registry.dispatch("generate_xlsx_from_xlsx", request("att_" + "0" * 32))
    assert not missing.success and "does not exist in this project" in missing.error
    foreign = await upload("other", workbook_bytes())
    assert not (await registry.dispatch("generate_xlsx_from_xlsx", request(foreign["id"]))).success
    assert (await db.fetchone("SELECT COUNT(*) AS n FROM office_deliveries"))["n"] == 0


async def test_import_obeys_prior_csv_and_table_denials():
    class Store:
        def lookup(self, name):
            return {"action": "deny"} if name == "generate_xlsx_from_csv" else None

    class Fake(ToolRegistry):
        async def dispatch(self, name, args):
            return ToolResult.ok()

    gate = type("Gater", (), {"request_consent": AsyncMock(return_value=True)})()
    core = AgentCore(llm=object(), registry=Fake(), permission_store=Store(), permission_gater=gate)
    assert not (await core._dispatch_tool({"name": "generate_xlsx_from_xlsx", "arguments": {}})).success
