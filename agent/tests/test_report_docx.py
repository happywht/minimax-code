from __future__ import annotations

import hashlib
import io
from io import BytesIO

import httpx
import pytest
from docx import Document

from minimax_code.attachments import AttachmentService
from minimax_code.storage.dao.projects import ProjectsDAO
from minimax_code.storage.db import AsyncDatabase


async def test_report_generation_registers_verified_bytes_without_changing_original(tmp_path):
    db = AsyncDatabase(tmp_path / "data" / "test.db")
    await db.connect()
    await db.migrate()
    root = tmp_path / "project"
    root.mkdir()
    original = root / "报告.docx"
    original.write_bytes(b"existing original")
    await ProjectsDAO(db).create(id="p", name="project", root_path=str(root))
    service = AttachmentService(db, db.path.parent)
    identity = await service.dao.service_id()
    result = await service.generate_docx("p", identity, "报告.docx", "资料报告", ["第一段", "来源：手工核对"])
    row = result["attachment"]
    assert row["parser_status"] == "parsed" and row["quality_status"] == "unverified"
    assert "第一段" in result["preview_text"]
    _, data = await service.content("p", row["id"])
    assert hashlib.sha256(data).hexdigest() == row["sha256"]
    assert [p.text for p in Document(BytesIO(data)).paragraphs] == ["资料报告", "第一段", "来源：手工核对"]
    assert original.read_bytes() == b"existing original"
    assert (await service.list("p"))["attachments"][0]["id"] == row["id"]
    await db.close()


async def test_report_rejects_invalid_inputs_before_publication(tmp_path):
    db = AsyncDatabase(tmp_path / "data" / "test.db")
    await db.connect()
    await db.migrate()
    await ProjectsDAO(db).create(id="p", name="project", root_path=str(tmp_path))
    service = AttachmentService(db, db.path.parent)
    identity = await service.dao.service_id()
    for filename, title, paragraphs in [("../bad.docx", "title", ["text"]),
                                        ("ok.docx", "title", ["bad\x00"]),
                                        ("ok.docx", "title", ["x" * 4097])]:
        with pytest.raises(ValueError):
            await service.generate_docx("p", identity, filename, title, paragraphs)
    assert not service.vault.exists()
    assert (await service.list("p"))["attachments"] == []
    await db.close()


async def test_unavailable_preview_does_not_register_document(tmp_path, monkeypatch):
    from minimax_code import attachment_parser

    db = AsyncDatabase(tmp_path / "data" / "test.db")
    await db.connect()
    await db.migrate()
    await ProjectsDAO(db).create(id="p", name="project", root_path=str(tmp_path))
    service = AttachmentService(db, db.path.parent)
    identity = await service.dao.service_id()

    async def unavailable(_data):
        return {"parser_status": "parser_unavailable"}

    monkeypatch.setattr(attachment_parser, "inspect_docx", unavailable)
    with pytest.raises(ValueError, match="could not be parsed"):
        await service.generate_docx("p", identity, "report.docx", "Report", ["Text"])
    assert (await service.list("p"))["attachments"] == []
    assert list(service.vault.iterdir()) == []
    await db.close()


async def test_authenticated_rpc_and_download(tmp_path, monkeypatch):
    from minimax_code import attachment_http
    from minimax_code.config import Config
    from minimax_code.http_server import build_app
    from minimax_code.ipc import handlers_attachments
    from minimax_code.ipc.server import IPCServer

    db = AsyncDatabase(tmp_path / "data" / "test.db")
    await db.connect()
    await db.migrate()
    await ProjectsDAO(db).create(id="p", name="project", root_path=str(tmp_path))
    service = AttachmentService(db, db.path.parent)
    identity = await service.dao.service_id()

    async def factory():
        return service

    monkeypatch.setattr(handlers_attachments, "get_attachment_service", factory)
    monkeypatch.setattr(attachment_http, "get_attachment_service", factory)
    server = IPCServer(config=Config.from_env(), stdin=io.StringIO(), stdout=io.StringIO())
    handlers_attachments.register_attachment_handlers(server)
    app = build_app(server, version="report-test")
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://testserver") as client:
        client.headers["Authorization"] = f"Bearer {app.state.auth.token}"
        payload = {"project_id": "p", "expected_service_id": identity, "filename": "report.docx",
                   "title": "Report", "paragraphs": ["Source: example", "Conclusion"]}
        response = await client.post("/rpc", json={"jsonrpc": "2.0", "id": 1,
                                                    "method": "attachment.generate_docx", "params": payload})
        result = response.json()["result"]
        row = result["attachment"]
        download = await client.get(f"/attachments/{row['id']}/content?project_id=p")
        assert download.status_code == 200
        assert download.headers["x-attachment-sha256"] == row["sha256"]
        assert [p.text for p in Document(BytesIO(download.content)).paragraphs] == ["Report", *payload["paragraphs"]]
        payload["unexpected"] = True
        response = await client.post("/rpc", json={"jsonrpc": "2.0", "id": 2,
                                                    "method": "attachment.generate_docx", "params": payload})
        assert response.json()["error"]["code"] == -32602
    server.listeners.clear()
    await db.close()


async def test_agent_tool_is_bound_to_session_and_plan_cannot_write(tmp_path):
    from minimax_code.agent.tools import ToolRegistry
    from minimax_code.agent.tools.generate_report_docx import register_report_docx_tool
    from minimax_code.execution import ExecutionPolicy, execution_scope
    from minimax_code.storage.dao.sessions import SessionsDAO

    db = AsyncDatabase(tmp_path / "data" / "test.db")
    await db.connect()
    await db.migrate()
    await ProjectsDAO(db).create(id="p", name="project", root_path=str(tmp_path))
    await SessionsDAO(db).create(id="s", project_id="p")
    registry = ToolRegistry()
    register_report_docx_tool(registry, db, "s", "p")
    args = {"filename": "report.docx", "title": "Report", "paragraphs": ["Evidence"]}
    with execution_scope(ExecutionPolicy(read_only=True)):
        assert not (await registry.dispatch("generate_report_docx", args)).success
    assert not (await registry.dispatch("generate_report_docx", {**args, "project_id": "p"})).success
    result = await registry.dispatch("generate_report_docx", args)
    assert result.success and result.output["quality_status"] == "unverified"
    await ProjectsDAO(db).create(id="other", name="other", root_path=str(tmp_path))
    await SessionsDAO(db).update("s", project_id="other")
    assert not (await registry.dispatch("generate_report_docx", args)).success
    assert len((await AttachmentService(db, db.path.parent).list("p"))["attachments"]) == 1
    await db.close()


async def test_session_rebind_during_generation_cannot_commit(tmp_path, monkeypatch):
    from minimax_code import attachment_parser
    from minimax_code.agent.tools import ToolRegistry
    from minimax_code.agent.tools.generate_report_docx import register_report_docx_tool
    from minimax_code.storage.dao.sessions import SessionsDAO

    db = AsyncDatabase(tmp_path / "data" / "test.db")
    await db.connect()
    await db.migrate()
    await ProjectsDAO(db).create(id="p", name="project", root_path=str(tmp_path))
    await ProjectsDAO(db).create(id="other", name="other", root_path=str(tmp_path))
    sessions = SessionsDAO(db)
    await sessions.create(id="s", project_id="p")
    original = attachment_parser.inspect_docx

    async def rebind(data):
        result = await original(data)
        await sessions.update("s", project_id="other")
        return result

    monkeypatch.setattr(attachment_parser, "inspect_docx", rebind)
    registry = ToolRegistry()
    register_report_docx_tool(registry, db, "s", "p")
    result = await registry.dispatch("generate_report_docx", {"filename": "report.docx", "title": "Report", "paragraphs": ["text"]})
    assert not result.success
    assert (await AttachmentService(db, db.path.parent).list("p"))["attachments"] == []
    await db.close()


@pytest.mark.parametrize(("rules", "has_gater", "consent", "allowed"), [
    ({}, False, False, False),
    ({}, True, False, False),
    ({}, True, True, True),
    ({"write_file": "deny", "generate_report_docx": "allow"}, True, True, False),
    ({"write_file": "ask", "generate_report_docx": "allow"}, True, False, False),
    ({"write_file": "allow", "generate_report_docx": "deny"}, True, True, False),
    ({"write_file": "allow"}, False, False, True),
])
async def test_agent_report_permission_combines_existing_write_rules(rules, has_gater, consent, allowed):
    from unittest.mock import AsyncMock

    from minimax_code.agent.core import AgentCore
    from minimax_code.agent.tools import ToolRegistry
    from minimax_code.agent.tools.base import Tool, ToolResult

    class Report(Tool):
        name = "generate_report_docx"
        parameters = {"type": "object", "properties": {}}

        async def run(self):
            return ToolResult.ok("created")

    class Store:
        def lookup(self, name):
            return {"action": rules[name]} if name in rules else None

    registry = ToolRegistry()
    registry.register(Report())
    gater = type("Gater", (), {"request_consent": AsyncMock(return_value=consent)})() if has_gater else None
    core = AgentCore(llm=object(), registry=registry, permission_store=Store(), permission_gater=gater)
    result = await core._dispatch_tool({"name": "generate_report_docx", "arguments": {}})
    assert result.success is allowed
