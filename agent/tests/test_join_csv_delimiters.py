from __future__ import annotations

from io import BytesIO

import openpyxl
import pytest

from minimax_code.agent.tools import ToolRegistry
from minimax_code.agent.tools.join_spreadsheets_xlsx import register_join_spreadsheet_tool
from minimax_code.attachments import AttachmentService
from minimax_code.execution import ExecutionPolicy, execution_scope
from minimax_code.office_deliveries import OfficeService
from minimax_code.storage.dao.projects import ProjectsDAO
from minimax_code.storage.dao.sessions import SessionsDAO
from minimax_code.storage.db import AsyncDatabase


@pytest.fixture
async def env(tmp_path):
    db = AsyncDatabase(tmp_path / "data" / "test.db")
    await db.connect()
    await db.migrate()
    projects = ProjectsDAO(db)
    await projects.create(id="p", name="P", root_path=str(tmp_path))
    await projects.create(id="other", name="Other", root_path=str(tmp_path))
    await SessionsDAO(db).create(id="s", project_id="p")
    service = AttachmentService(db, db.path.parent)
    identity = await service.dao.service_id()

    async def upload(name, raw, project="p"):
        async def chunks():
            yield raw
        return await service.upload(project, name, identity, chunks())

    registry = ToolRegistry()
    register_join_spreadsheet_tool(registry, db, "s", "p")
    yield db, service, upload, registry
    await db.close()


def source(attachment_id, columns, **extra):
    return {"attachment_id": attachment_id, "columns": columns, "key": "ID", **extra}


def request(left, right):
    return {"left": left, "right": right, "filename": "result.xlsx",
            "request_key": "join-key", "sheet_name": "Data"}


ID = {"source": "id", "name": "ID", "kind": "text"}
VALUE = {"source": "value", "name": "Value", "kind": "text"}


async def test_join_independent_delimiters_preserves_originals_and_receipt(env):
    db, service, upload, registry = env
    left_bytes = b'id;left\n"A;1";first\n'
    right_bytes = b'id\tvalue\n"A;1"\t"x\ty"\n'
    left = await upload("left.csv", left_bytes)
    right = await upload("right.csv", right_bytes)
    args = request(source(left["id"], [ID], delimiter="semicolon"),
                   source(right["id"], [ID, VALUE], delimiter="tab"))
    with execution_scope(ExecutionPolicy(read_only=True)):
        assert not (await registry.dispatch("join_spreadsheets_xlsx", args)).success
    result = await registry.dispatch("join_spreadsheets_xlsx", args)
    assert result.success, result.error
    _, data = await service.content("p", result.output["attachment_id"])
    sheet = openpyxl.load_workbook(BytesIO(data)).active
    assert [sheet["A2"].value, sheet["B2"].value] == ["A;1", "x\ty"]
    receipt = await OfficeService(db, db.path.parent).get("p", result.output["delivery_id"])
    sources = receipt["delivery"]["validation_report"]["processing"]["sources"]
    assert [item["csv_delimiter"] for item in sources] == ["semicolon", "tab"]
    assert (await service.content("p", left["id"]))[1] == left_bytes
    assert (await service.content("p", right["id"]))[1] == right_bytes
    retry = await registry.dispatch("join_spreadsheets_xlsx", args)
    assert retry.success and retry.output["delivery_id"] == result.output["delivery_id"]


async def test_join_delimiter_change_conflicts_and_foreign_source_denied(env):
    _, _, upload, registry = env
    left = await upload("left.csv", b"id\nA\n")
    right = await upload("right.csv", b"id,value\nA,one\n")
    args = request(source(left["id"], [ID]), source(right["id"], [ID, VALUE]))
    first = await registry.dispatch("join_spreadsheets_xlsx", args)
    assert first.success, first.error
    changed = {**args, "left": {**args["left"], "delimiter": "semicolon"}}
    assert not (await registry.dispatch("join_spreadsheets_xlsx", changed)).success
    foreign = await upload("foreign.csv", b"id\nA\n", project="other")
    changed["left"] = source(foreign["id"], [ID], delimiter="tab")
    assert not (await registry.dispatch("join_spreadsheets_xlsx", changed)).success


async def test_join_csv_and_xlsx_sources(env):
    _, service, upload, registry = env
    left_bytes = b'id;other\nA;data\n'
    left = await upload("left.csv", left_bytes)
    workbook = openpyxl.Workbook()
    workbook.active.append(["id", "value"])
    workbook.active.append(["A", "found"])
    output = BytesIO()
    workbook.save(output)
    right_bytes = output.getvalue()
    right = await upload("right.xlsx", right_bytes)
    args = request(source(left["id"], [ID], delimiter="semicolon"),
                   source(right["id"], [ID, VALUE], source_sheet="Sheet"))
    result = await registry.dispatch("join_spreadsheets_xlsx", args)
    assert result.success, result.error
    _, data = await service.content("p", result.output["attachment_id"])
    assert openpyxl.load_workbook(BytesIO(data)).active["B2"].value == "found"
    assert (await service.content("p", left["id"]))[1] == left_bytes
    assert (await service.content("p", right["id"]))[1] == right_bytes


async def test_join_rejects_malformed_csv_and_delimiter_for_xlsx(env):
    db, _, upload, registry = env
    malformed = await upload("left.csv", b'id;value\n"unterminated;item')
    right = await upload("right.csv", b"id,value\nA,one\n")
    args = request(source(malformed["id"], [ID], delimiter="semicolon"),
                   source(right["id"], [ID, VALUE]))
    result = await registry.dispatch("join_spreadsheets_xlsx", args)
    assert not result.success and "malformed" in result.error
    workbook = openpyxl.Workbook()
    workbook.active.append(["id"])
    workbook.active.append(["A"])
    output = BytesIO()
    workbook.save(output)
    excel = await upload("left.xlsx", output.getvalue())
    args["left"] = source(excel["id"], [ID], source_sheet="Sheet", delimiter="comma")
    result = await registry.dispatch("join_spreadsheets_xlsx", args)
    assert not result.success and "cannot specify" in result.error
    assert (await db.fetchone("SELECT COUNT(*) AS n FROM office_deliveries"))["n"] == 0
