"""Project-bound CSV original conversion without network or model calls."""
from __future__ import annotations

from io import BytesIO
from unittest.mock import AsyncMock

import openpyxl
import pytest

from minimax_code.agent.core import AgentCore
from minimax_code.agent.tools import ToolRegistry
from minimax_code.agent.tools.base import ToolResult
from minimax_code.agent.tools.generate_xlsx_from_csv import csv_to_spec, register_csv_xlsx_tool
from minimax_code.attachments import AttachmentService
from minimax_code.execution import ExecutionPolicy, execution_scope
from minimax_code.office_deliveries import OfficeService
from minimax_code.storage.dao.projects import ProjectsDAO
from minimax_code.storage.dao.sessions import SessionsDAO
from minimax_code.storage.db import AsyncDatabase


@pytest.fixture
async def env(tmp_path):
    db = AsyncDatabase(tmp_path / "data" / "test.db")
    await db.connect()
    await db.migrate()
    projects = ProjectsDAO(db)
    await projects.create(id="p", name="P", root_path=str(tmp_path))
    await projects.create(id="other", name="Other", root_path=str(tmp_path))
    sessions = SessionsDAO(db)
    await sessions.create(id="s", project_id="p")
    service = AttachmentService(db, db.path.parent)
    identity = await service.dao.service_id()

    async def upload(project, raw):
        async def chunks():
            yield raw
        return await service.upload(project, "data.csv", identity, chunks())

    registry = ToolRegistry()
    register_csv_xlsx_tool(registry, db, "s", "p")
    yield db, sessions, service, upload, registry
    await db.close()


def request(attachment_id):
    return {"attachment_id": attachment_id, "filename": "result.xlsx", "request_key": "key1",
            "sheet_name": "结果", "columns": [{"source": "数量", "name": "金额", "kind": "number"},
                                                     {"source": "名称", "name": "项目", "kind": "text"}]}


async def test_full_original_beyond_preview_reorder_summary_and_retry(env):
    db, _, service, upload, registry = env
    csv_data = ("\ufeff名称,数量\n" + "".join(f"项目{i}-" + "甲" * 90 + f",{i}\n" for i in range(800))).encode()
    row = await upload("p", csv_data)
    assert row["parser_status"] == "truncated"
    data = request(row["id"])
    data["summaries"] = [{"column": "金额", "operation": "sum"}]
    with execution_scope(ExecutionPolicy(read_only=True)):
        assert not (await registry.dispatch("generate_xlsx_from_csv", data)).success
    result = await registry.dispatch("generate_xlsx_from_csv", data)
    assert result.success, result.error
    assert result.output["source_sha256"] == row["sha256"]
    office = OfficeService(db, db.path.parent)
    _, xlsx = await office.originals.content("p", result.output["attachment_id"])
    sheet = openpyxl.load_workbook(BytesIO(xlsx), data_only=False).active
    assert sheet["A1"].value == "金额" and sheet["B1"].value == "项目"
    assert sheet["A801"].value == 799 and sheet["B801"].value.startswith("项目799")
    assert sheet["D802"].value == sum(range(800))
    assert (await service.content("p", row["id"]))[1] == csv_data
    repeat = await registry.dispatch("generate_xlsx_from_csv", data)
    assert repeat.success and repeat.output["delivery_id"] == result.output["delivery_id"]


async def test_csv_chart_and_project_isolation(env):
    db, sessions, _, upload, registry = env
    missing = await registry.dispatch("generate_xlsx_from_csv", request("att_" + "0" * 32))
    assert not missing.success and "does not exist in this project" in missing.error
    foreign = await upload("other", "名称,数量\n甲,3".encode())
    assert not (await registry.dispatch("generate_xlsx_from_csv", request(foreign["id"]))).success
    row = await upload("p", "名称,数量\n甲,3\n乙,5".encode())
    data = request(row["id"])
    data["chart"] = {"kind": "bar", "category_column": "项目", "value_column": "金额", "title": "金额"}
    result = await registry.dispatch("generate_xlsx_from_csv", data)
    assert result.success, result.error
    assert result.output["validation_scope"] == "xlsx_chart_values_v1"
    attachment = AttachmentService(db, db.path.parent)
    _, blob = await attachment.content("p", result.output["attachment_id"])
    assert len(openpyxl.load_workbook(BytesIO(blob)).active._charts) == 1
    await sessions.update("s", project_id="other")
    assert not (await registry.dispatch("generate_xlsx_from_csv", data)).success


@pytest.mark.parametrize("delimiter,character", [("semicolon", ";"), ("tab", "\t")])
async def test_explicit_delimiter_creates_recorded_delivery_without_changing_original(env, delimiter, character):
    db, _, service, upload, registry = env
    original = f'名称{character}数量\n"甲{character}乙"{character}3\n'.encode()
    row = await upload("p", original)
    data = {**request(row["id"]), "delimiter": delimiter}
    with execution_scope(ExecutionPolicy(read_only=True)):
        assert not (await registry.dispatch("generate_xlsx_from_csv", data)).success
    result = await registry.dispatch("generate_xlsx_from_csv", data)
    assert result.success, result.error
    office = OfficeService(db, db.path.parent)
    _, blob = await service.content("p", result.output["attachment_id"])
    sheet = openpyxl.load_workbook(BytesIO(blob)).active
    assert [sheet["A2"].value, sheet["B2"].value] == [3, f"甲{character}乙"]
    receipt = await office.get("p", result.output["delivery_id"])
    assert receipt["delivery"]["validation_report"]["processing"]["csv_delimiter"] == delimiter
    assert (await service.content("p", row["id"]))[1] == original
    retry = await registry.dispatch("generate_xlsx_from_csv", data)
    assert retry.success and retry.output["delivery_id"] == result.output["delivery_id"]


async def test_delimiter_participates_in_request_conflicts(env):
    _, _, _, upload, registry = env
    row = await upload("p", b"name\nitem\n")
    data = {"attachment_id": row["id"], "filename": "result.xlsx", "request_key": "same",
            "sheet_name": "Data", "columns": [{"source": "name", "name": "Name", "kind": "text"}]}
    first = await registry.dispatch("generate_xlsx_from_csv", data)
    assert first.success, first.error
    conflict = await registry.dispatch("generate_xlsx_from_csv", {**data, "delimiter": "semicolon"})
    assert not conflict.success


async def test_invalid_explicit_delimiter_data_does_not_create_delivery(env):
    db, _, service, upload, registry = env
    original = b'name;amount\n"unterminated;3'
    row = await upload("p", original)
    args = {"attachment_id": row["id"], "filename": "bad.xlsx", "request_key": "bad-dialect",
            "sheet_name": "Data", "delimiter": "semicolon",
            "columns": [{"source": "name", "name": "Name", "kind": "text"}]}
    result = await registry.dispatch("generate_xlsx_from_csv", args)
    assert not result.success and "malformed" in result.error
    assert (await db.fetchone("SELECT COUNT(*) AS n FROM office_deliveries"))["n"] == 0
    assert (await service.content("p", row["id"]))[1] == original


async def test_conversion_error_does_not_create_result(env):
    db, _, service, upload, registry = env
    original = b"name,amount\nitem,not-a-number\n"
    row = await upload("p", original)
    data = {"attachment_id": row["id"], "filename": "bad.xlsx", "request_key": "bad",
            "sheet_name": "Data", "columns": [{"source": "amount", "name": "Amount", "kind": "number"}]}
    result = await registry.dispatch("generate_xlsx_from_csv", data)
    assert not result.success and "invalid number" in result.error
    assert (await db.fetchone("SELECT COUNT(*) AS n FROM office_deliveries"))["n"] == 0
    assert len((await service.list("p"))["attachments"]) == 1
    assert (await service.content("p", row["id"]))[1] == original


def test_csv_errors_and_type_semantics():
    columns = [{"source": "v", "name": "Value", "kind": "number"}]
    with pytest.raises(ValueError, match="delimiter"):
        csv_to_spec(b"v\n1", "Data", columns, delimiter="pipe")
    for raw in (b"\xff", b'v\n"unclosed', b"v,v\n1,2", b"v\n12,3", b"v\nNaN", b"v\n1e100"):
        with pytest.raises(ValueError):
            spec = csv_to_spec(raw, "Data", columns)
            from minimax_code.office_xlsx import XlsxSpec
            XlsxSpec.model_validate(spec)
    result = csv_to_spec(b"v\n\"\"\n1.25", "Data", columns)
    assert result["rows"] == [[None], [1.25]]
    assert csv_to_spec(b"v\n=true", "Data", [{"source": "v", "name": "Value", "kind": "text"}])["rows"] == [["=true"]]
    assert csv_to_spec(b"v\nfalse\nTRUE", "Data", [{"source": "v", "name": "Flag", "kind": "boolean"}])["rows"] == [[False], [True]]


async def test_csv_write_permissions_include_existing_table_rule():
    class Store:
        def lookup(self, name):
            return {"action": "deny"} if name == "generate_table_xlsx" else None

    class Fake(ToolRegistry):
        async def dispatch(self, name, args):
            return ToolResult.ok()

    call = {"name": "generate_xlsx_from_csv", "arguments": {}}
    core = AgentCore(llm=object(), registry=Fake(), permission_store=Store(),
                     permission_gater=type("Gater", (), {"request_consent": AsyncMock(return_value=True)})())
    assert not (await core._dispatch_tool(call)).success
