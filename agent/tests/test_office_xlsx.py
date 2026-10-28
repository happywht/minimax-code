"""Actual XLSX bytes, whole-file comparison and fixed-worker resource ownership."""
from __future__ import annotations

import asyncio
import io
import subprocess
import sys
import zipfile

import pytest
from pydantic import ValidationError

from minimax_code import office_process as process
from minimax_code.office_xlsx import XlsxSpec, filename, generate, verify


def spec(rows=None):
    return XlsxSpec.model_validate({
        "sheet_name": "清单", "columns": [{"name": "名称", "kind": "text"},
                                         {"name": "数量", "kind": "number"}],
        "rows": [["=SUM(1,2)", 123.45], ["#N/A", None], ["", 3]] if rows is None else rows,
    })


def mutate(data, part, transform):
    result = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(result, "w") as target:
        for entry in source.infolist():
            raw = source.read(entry.filename)
            target.writestr(entry, transform(raw) if entry.filename == part else raw)
    return result.getvalue()


def test_all_literal_prefixes_and_complete_roundtrip():
    from openpyxl import load_workbook

    values = ["=1+1", "+SUM(1,2)", "-1+2", "@SUM(1,2)", "#N/A", "=cmd|' /C calc'!A0",
              "https://example.invalid/", "中文<&>😀\n第二行", "0001234567890123456789"]
    request = spec([[value, index] for index, value in enumerate(values)])
    data = generate(request)
    report = verify(data, request)
    assert report["cell_count"] == 20 and report["business_quality"] == "unverified"
    assert report["render_status"] == "not_attempted"
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        assert [(book.active.cell(index + 2, 1).value, book.active.cell(index + 2, 1).data_type)
                for index in range(len(values))] == [(value, "s") for value in values]
    finally:
        book.close()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        sheet = archive.read("xl/worksheets/sheet1.xml")
        assert b"<f>" not in sheet and b"hyperlink" not in sheet
        assert not any("vba" in name.lower() or "externallink" in name.lower()
                       for name in archive.namelist())


def test_empty_text_is_explicitly_normalized_and_boolean_not_number():
    request = spec()
    assert request.rows[-1][0] is None
    verify(generate(request), request)
    boolean = XlsxSpec.model_validate({"sheet_name": "bool", "columns": [{"name": "b", "kind": "boolean"}],
                                       "rows": [[True], [False], [None]]})
    assert verify(generate(boolean), boolean)["data_rows"] == 3


def test_exact_cell_budget_and_input_byte_budget():
    columns = [{"name": f"column {index}", "kind": "text"} for index in range(20)]
    maximum = XlsxSpec.model_validate({"sheet_name": "data", "columns": columns,
                                      "rows": [[None] * 20 for _ in range(249)]})
    assert verify(generate(maximum), maximum)["cell_count"] == 5000
    with pytest.raises(ValueError, match="5000 cells"):
        XlsxSpec.model_validate({"sheet_name": "data", "columns": columns,
                                "rows": [[None] * 20 for _ in range(250)]})
    # All per-cell/row bounds pass; aggregate UTF-8 bytes independently refuse.
    with pytest.raises(ValueError, match="1 MiB"):
        XlsxSpec.model_validate({"sheet_name": "data", "columns": columns[:1],
                                "rows": [["汉" * 4096] for _ in range(100)]})
    with pytest.raises(ValueError):
        verify(b"x" * (2 * 1024 * 1024 + 1), maximum)


@pytest.mark.parametrize("value", [True, "3", float("nan"), float("inf"), -0.0,
                                   1234567890123456, 0.1234567890123456, 1e-13, 10 ** 1000])
def test_number_restrictions(value):
    with pytest.raises(ValueError):
        spec([["number", value]])


@pytest.mark.parametrize("change", [
    {"sheet_name": "bad/name"}, {"sheet_name": "'quote"}, {"sheet_name": "x" * 32},
    {"rows": [["short"]]}, {"rows": [["x\x00", 1]]}, {"rows": [["x" * 4097, 1]]},
    {"rows": [["x", 1]] * 1001}, {"unknown": True},
    {"columns": [{"name": "x", "kind": "formula"}]},
    {"columns": [{"name": "x", "kind": "text"}, {"name": "X", "kind": "text"}]},
])
def test_strict_spec_rejects_instead_of_truncating(change):
    payload = spec().model_dump()
    payload.update(change)
    with pytest.raises(ValidationError):
        XlsxSpec.model_validate(payload)


@pytest.mark.parametrize("name", ["../x.xlsx", "C:x.xlsx", "x.xlsm", "CON.xlsx", "bad\n.xlsx", " x.xlsx"])
def test_filename_is_display_basename(name):
    with pytest.raises(ValueError):
        filename(name)


@pytest.mark.parametrize("kind", ["tail", "formula", "external", "extra_part", "extra_cell"])
def test_saved_package_rejects_tampering_outside_preview_window(kind):
    request = spec([["tail" if index == 999 else "row", index] for index in range(1000)])
    data = generate(request)
    if kind == "tail":
        data = mutate(data, "xl/worksheets/sheet1.xml", lambda raw: raw.replace(b">tail<", b">wrong<"))
    elif kind == "formula":
        data = mutate(data, "xl/worksheets/sheet1.xml", lambda raw: raw.replace(b"<v>999</v>", b"<f>1+1</f><v>999</v>"))
    elif kind == "external":
        data = mutate(data, "xl/_rels/workbook.xml.rels", lambda raw: raw.replace(b"</Relationships>",
            b'<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Id="evil" Target="http://example.invalid/" TargetMode="External"/></Relationships>'))
    elif kind == "extra_cell":
        data = mutate(data, "xl/worksheets/sheet1.xml", lambda raw: raw.replace(b"</sheetData>",
            b'<row r="1002"><c r="A1002" t="inlineStr"><is><t>extra</t></is></c></row></sheetData>'))
    else:
        output = io.BytesIO(data)
        with zipfile.ZipFile(output, "a") as archive:
            archive.writestr("xl/vbaProject.bin", b"never execute")
        data = output.getvalue()
    with pytest.raises(ValueError):
        verify(data, request)


async def test_real_limited_worker_environment_and_report(monkeypatch):
    monkeypatch.setenv("MINIMAX_API_KEY", "must-not-reach-worker")
    original, children = subprocess.Popen, []

    def start(*args, **kwargs):
        assert kwargs["env"].get("MINIMAX_API_KEY") is None
        assert args[0][1:4] == ["-I", "-B", "-S"]
        child = original(*args, **kwargs)
        children.append(child)
        return child

    monkeypatch.setattr(process.subprocess, "Popen", start)
    data, report = await process.generate_xlsx(spec())
    assert report == verify(data, spec())
    assert len(children) == 1 and children[0].poll() == 0


@pytest.mark.parametrize("mode", ["timeout", "cancel", "stdout", "stderr", "invalid"])
async def test_worker_failure_and_repeated_cancellation_reap_child(monkeypatch, mode):
    original, children = subprocess.Popen, []

    def start(*args, **kwargs):
        child = original(*args, **kwargs)
        children.append(child)
        return child

    command = "import time;time.sleep(30)"
    if mode in {"stdout", "stderr"}:
        command = f"import sys,time;sys.{mode}.write('x'*100000);sys.{mode}.flush();time.sleep(30)"
        monkeypatch.setattr(process, "STDOUT_LIMIT" if mode == "stdout" else "STDERR_LIMIT", 1000)
    elif mode == "invalid":
        command = "print('not-json')"
    monkeypatch.setattr(process, "_command", lambda: [sys.executable, "-I", "-B", "-S", "-c", command])
    monkeypatch.setattr(process.subprocess, "Popen", start)
    monkeypatch.setattr(process, "WALL_SECONDS", .2)
    task = asyncio.create_task(process.generate_xlsx(spec()))
    if mode == "cancel":
        while not children:
            await asyncio.sleep(.01)
        task.cancel()
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
    else:
        with pytest.raises(process.OfficeProcessError) as error:
            await task
        assert "not-json" not in str(error.value)
    assert children and children[0].poll() is not None
    assert process._GATE.acquire(blocking=False)
    process._GATE.release()


@pytest.mark.parametrize("kind", ["cpu", "memory"])
async def test_actual_resource_gate_reaps_office_worker(monkeypatch, tmp_path, kind):
    # Use the actual trusted bootstrap and OS gates, changing only its post-gate
    # generation operation in a test-only copy. Production exposes no probe mode.
    import time
    from pathlib import Path

    original_command = process._command
    original = Path(original_command()[4])
    source = original.read_text(encoding="utf8").replace("__file__", repr(str(original)))
    expression = "exec('while True: pass')" if kind == "cpu" else "data = bytearray(800 * 1024 * 1024)"
    source = source.replace("data = generate(spec)", expression)
    probe = tmp_path / "resource_probe.py"
    probe.write_text(source, encoding="utf8")
    command = original_command()
    command[4] = str(probe)
    monkeypatch.setattr(process, "_command", lambda: command)
    monkeypatch.setattr(process, "WALL_SECONDS", 40)
    popen, children = subprocess.Popen, []

    def launched(*args, **kwargs):
        child = popen(*args, **kwargs)
        children.append(child)
        return child

    monkeypatch.setattr(process.subprocess, "Popen", launched)
    started = time.monotonic()
    with pytest.raises(process.OfficeProcessError):
        await process.generate_xlsx(spec())
    assert len(children) == 1 and children[0].poll() is not None
    assert time.monotonic() - started < (40 if kind == "cpu" else 14)
    monkeypatch.setattr(process, "_command", original_command)
    assert (await process.generate_xlsx(spec()))[1]["structure_status"] == "passed"


@pytest.mark.parametrize("kind", ["limits", "missing_dependency"])
async def test_actual_worker_fails_closed_without_gate_or_dependency(monkeypatch, kind):
    command = process._command()
    if kind == "limits":
        command[5] = "0"  # Invalid parent identity, before third parties or stdin.
    else:
        command[-2:] = ["/does-not-exist", "/does-not-exist"]
    monkeypatch.setattr(process, "_command", lambda: command)
    with pytest.raises(process.OfficeProcessError) as caught:
        await process.generate_xlsx(spec())
    assert caught.value.reason == ("limits_unavailable" if kind == "limits" else "missing_dependency")


@pytest.mark.parametrize("kind", ["hash", "input", "rows", "unknown"])
async def test_parent_rejects_inconsistent_worker_report(monkeypatch, kind):
    import base64
    import json

    data = generate(spec())
    report = verify(data, spec())
    report[{"hash": "file_sha256", "input": "input_sha256", "rows": "data_rows", "unknown": "unknown"}[kind]] = (
        "0" * 64 if kind in {"hash", "input"} else 999 if kind == "rows" else True)
    output = json.dumps({"data_base64": base64.b64encode(data).decode(), "report": report})
    monkeypatch.setattr(process, "_command", lambda: [sys.executable, "-I", "-B", "-S", "-c", f"print({output!r})"])
    with pytest.raises(process.OfficeProcessError, match="invalid_worker_output"):
        await process.generate_xlsx(spec())
