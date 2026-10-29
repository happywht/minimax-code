"""Partial and complete report pair behavior on real attachment storage."""
from __future__ import annotations

from unittest.mock import AsyncMock

import pytest

from minimax_code.agent.core import AgentCore
from minimax_code.agent.tools import ToolRegistry
from minimax_code.agent.tools.base import ToolResult
from minimax_code.agent.tools.generate_report_docx import register_report_docx_tool
from minimax_code.attachments import AttachmentService
from minimax_code.execution import ExecutionPolicy, execution_scope
from minimax_code.storage.dao.projects import ProjectsDAO
from minimax_code.storage.dao.sessions import SessionsDAO
from minimax_code.storage.db import AsyncDatabase


@pytest.fixture
async def env(tmp_path):
    db = AsyncDatabase(tmp_path / "data" / "test.db")
    await db.connect()
    await db.migrate()
    await ProjectsDAO(db).create(id="p", name="Project", root_path=str(tmp_path))
    await ProjectsDAO(db).create(id="other", name="Other", root_path=str(tmp_path))
    sessions = SessionsDAO(db)
    await sessions.create(id="s", project_id="p")
    registry = ToolRegistry()
    register_report_docx_tool(registry, db, "s", "p")
    yield db, sessions, registry, tmp_path
    await db.close()


def args():
    return {"filename": "报告", "title": "中文报告", "sections": [{"heading": "一、结论", "blocks": [
        {"kind": "paragraph", "text": "正文"}, {"kind": "table", "rows": [["项目", "数量"], ["钢材", "2"]]}]}]}


async def test_pair_returns_saved_docx_when_pdf_fails_and_retry_only_pdf(env, monkeypatch):
    from minimax_code import report_pdf

    db, _, registry, root = env
    original = root / "报告.docx"
    original.write_bytes(b"untouched")
    monkeypatch.setattr(report_pdf, "_font_paths", lambda: [])
    result = await registry.dispatch("generate_report_pair", args())
    assert not result.success and result.output["status"] == "partial"
    docx, pdf = result.output["results"]
    assert docx["status"] == "saved" and pdf["status"] == "unconfirmed"
    assert "MINIMAX_CODE_PDF_FONT" in pdf["error"]
    assert original.read_bytes() == b"untouched"
    service = AttachmentService(db, db.path.parent)
    assert len((await service.list("p"))["attachments"]) == 1
    again = await registry.dispatch("generate_report_pair", {**args(), "formats": ["pdf"]})
    assert not again.success and again.output["status"] == "unconfirmed"
    assert len((await service.list("p"))["attachments"]) == 1


async def test_complete_pair_contains_same_chinese_text(env):
    from io import BytesIO

    from docx import Document
    from pypdf import PdfReader

    from minimax_code.report_pdf import _font_paths

    if not any(path.is_file() for path in _font_paths()):
        pytest.skip("A compatible Chinese font is required")
    db, _, registry, _ = env
    result = await registry.dispatch("generate_report_pair", args())
    assert result.success and result.output["status"] == "complete"
    service = AttachmentService(db, db.path.parent)
    files = {item["format"]: (await service.content("p", item["attachment_id"]))[1]
             for item in result.output["results"]}
    doc = Document(BytesIO(files["docx"]))
    assert "正文" in [p.text for p in doc.paragraphs] and doc.tables[0].cell(1, 0).text == "钢材"
    text = "".join(page.extract_text() for page in PdfReader(BytesIO(files["pdf"])).pages)
    assert "正文" in text and "钢材" in text


async def test_binding_change_stops_next_format(env, monkeypatch):
    _, sessions, registry, _ = env
    tool = registry.get("generate_report_pair")
    original = tool.service.generate_docx
    pdf = AsyncMock()

    async def rebind(*args, **kwargs):
        result = await original(*args, **kwargs)
        await sessions.update("s", project_id="other")
        return result

    monkeypatch.setattr(tool.service, "generate_docx", rebind)
    monkeypatch.setattr(tool.service, "generate_pdf", pdf)
    result = await registry.dispatch("generate_report_pair", args())
    assert not result.success and "original project's attachments" in result.error
    pdf.assert_not_awaited()


async def test_plan_and_existing_pdf_deny_rules(env):
    _, _, registry, _ = env
    with execution_scope(ExecutionPolicy(read_only=True)):
        assert not (await registry.dispatch("generate_report_pair", args())).success

    class Store:
        def lookup(self, name):
            return {"action": "deny"} if name == "generate_report_pdf" else None

    class Fake(ToolRegistry):
        async def dispatch(self, name, args):
            return ToolResult.ok()

    gater = type("Gater", (), {"request_consent": AsyncMock(return_value=True)})()
    core = AgentCore(llm=object(), registry=Fake(), permission_store=Store(), permission_gater=gater)
    assert not (await core._dispatch_tool({"name": "generate_report_pair", "arguments": args()})).success

    class OnlyDocxAllowed:
        def lookup(self, name):
            return {"action": "allow"} if name == "generate_report_docx" else None

    no_gater = AgentCore(llm=object(), registry=Fake(), permission_store=OnlyDocxAllowed())
    assert not (await no_gater._dispatch_tool({"name": "generate_report_pair", "arguments": args()})).success
    single = {**args(), "formats": ["docx"]}
    assert (await no_gater._dispatch_tool({"name": "generate_report_pair", "arguments": single})).success
