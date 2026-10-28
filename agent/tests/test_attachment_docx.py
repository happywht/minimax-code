"""Real DOCX packages, hostile ZIP/XML and actual bounded child processes."""
from __future__ import annotations

import asyncio
import hashlib
import io
import os
import struct
import subprocess
import sys
import sysconfig
import time
import zipfile
from pathlib import Path

import pytest

from minimax_code import attachment_parser as parser
from minimax_code.attachments import AttachmentService
from minimax_code.storage.dao.projects import ProjectsDAO
from minimax_code.storage.db import AsyncDatabase

if os.environ.get("MINIMAX_REQUIRE_ATTACHMENT_EXTRA") == "1":
    import defusedxml  # noqa: F401 - required dependency gate for feature CI
    import docx
else:
    docx = pytest.importorskip("docx")
    pytest.importorskip("defusedxml")


def document(text="中文正文"):
    value = docx.Document()
    value.add_paragraph(text)
    return value


def encoded(value):
    stream = io.BytesIO()
    value.save(stream)
    return stream.getvalue()


def altered(data, changes=None, extra=()):
    result = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(result, "w", compression=8) as target:
        for row in source.infolist():
            value = source.read(row)
            if changes and row.filename in changes:
                change = changes[row.filename]
                value = change(value) if callable(change) else change
            target.writestr(row.filename, value)
        for name, value in extra:
            target.writestr(name, value)
    return result.getvalue()


async def test_actual_child_chinese_body_order_nested_tables_and_omissions():
    value = document("开头")
    value.sections[0].header.paragraphs[0].text = "不能出现的页眉"
    value.sections[0].footer.paragraphs[0].text = "不能出现的页脚"
    table = value.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "第一格"
    table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "嵌套表"
    table.cell(0, 1).text = "第二格"
    value.add_paragraph("结尾")
    # Unsupported stories/revisions/textboxes must not become apparent body text.
    from docx.oxml import parse_xml
    value._element.body.insert(0, parse_xml(
        '<w:ins xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:p><w:r><w:t>不能出现的修订</w:t></w:r></w:p></w:ins>'))
    p = value.add_paragraph()
    p._p.append(parse_xml('<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                         '<w:pict><w:txbxContent><w:p><w:r><w:t>不能出现的文本框</w:t>'
                         '</w:r></w:p></w:txbxContent></w:pict></w:r>'))
    result = await parser.inspect_docx(encoded(value))
    assert result["parser_status"] == "parsed", result
    preview = result["preview_text"]
    assert [preview.index(word) for word in ("开头", "第一格", "嵌套表", "第二格", "结尾")] == sorted(
        preview.index(word) for word in ("开头", "第一格", "嵌套表", "第二格", "结尾"))
    assert "[body/2/row1/cell1/2/row1/cell1/1] 嵌套表" in preview
    assert "不能出现" not in preview
    assert result["parser_scope"] == "docx_body_paragraphs_tables" and result["parser_version"] == 2


async def test_actual_child_empty_long_and_nested_limit():
    value = docx.Document()
    value.sections[0].header.paragraphs[0].text = "header-only"
    result = await parser.inspect_docx(encoded(value))
    assert result["parser_status"] == "empty" and result["parser_reason"] == "no_extracted_body_text_no_ocr"
    assert result["preview_text"] == ""
    result = await parser.inspect_docx(encoded(document("长正文" * 40_000)))
    assert result["parser_status"] == "truncated" and result["preview_truncated"] == 1
    assert len(result["preview_text"].encode()) <= 65536 and len(result["preview_text"].splitlines()) <= 2000
    value = docx.Document()
    container = value
    for _ in range(9):
        container = container.add_table(rows=1, cols=1).cell(0, 0)
    container.text = "too deep"
    result = await parser.inspect_docx(encoded(value))
    assert result["parser_reason"] == "package_limit" and result["preview_text"] == ""


@pytest.mark.parametrize("kind", ["duplicate", "traversal", "backslash", "case-alias", "encrypted",
                                  "zip64", "multivolume", "compression", "crc", "size", "bomb",
                                  "macro", "ole", "entity-utf16", "dtd-unused", "broken-relation",
                                  "external-package", "missing-content-type", "missing-root-relation",
                                  "actual-bomb", "total-bomb", "member-count"])
async def test_hostile_packages_rejected_in_actual_child(kind):
    data = encoded(document())
    if kind == "actual-bomb":
        raw = bytearray(altered(data, extra=[("word/actual-bomb.xml", b"x" * (8 * 1024 * 1024 + 1))]))
        cursor = raw.index(b"PK\x01\x02")
        while raw[cursor:cursor + 4] == b"PK\x01\x02":
            length, extra, note = struct.unpack_from("<3H", raw, cursor + 28)
            if raw[cursor + 46:cursor + 46 + length] == b"word/actual-bomb.xml":
                local = struct.unpack_from("<I", raw, cursor + 42)[0]
                struct.pack_into("<I", raw, cursor + 24, 1)
                struct.pack_into("<I", raw, local + 22, 1)
                break
            cursor += 46 + length + extra + note
        data = bytes(raw)
    elif kind == "total-bomb":
        data = altered(data, extra=[(f"word/bomb{n}.xml", b"x" * (7 * 1024 * 1024)) for n in range(5)])
    elif kind == "member-count":
        data = altered(data, extra=[(f"word/extra{n}.xml", b"<x/>") for n in range(2049)])
    elif kind in {"duplicate", "traversal", "backslash", "case-alias", "bomb", "macro", "ole", "dtd-unused"}:
        name, value = {
            "duplicate": ("word/document.xml", b"<x/>"),
            "traversal": ("../escape.xml", b"<x/>"),
            "backslash": ("word/evil.xml", b"<x/>"),
            "case-alias": ("WORD/document.xml", b"<x/>"),
            "bomb": ("word/bomb.xml", b"x" * (8 * 1024 * 1024 + 1)),
            "macro": ("word/vbaProject.bin", b"macro"),
            "ole": ("word/embeddings/item.bin", b"object"),
            "dtd-unused": ("word/unused.xml", b'<!DOCTYPE x SYSTEM "http://127.0.0.1:9/private"><x/>'),
        }[kind]
        data = altered(data, extra=[(name, value)])
        if kind == "backslash":
            data = data.replace(b"word/evil.xml", b"word\\evil.xml")
    elif kind == "entity-utf16":
        xml = '<?xml version="1.0" encoding="utf-16"?><!DOCTYPE x [<!ENTITY e SYSTEM "file:///private">]><x>&e;</x>'
        data = altered(data, {"word/document.xml": xml.encode("utf16")})
    elif kind in {"broken-relation", "external-package"}:
        content = ('<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="bad" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
                   + ('package" Target="http://127.0.0.1:9/private" TargetMode="External"' if kind == "external-package" else 'image" Target="missing.png"')
                   + '/></Relationships>')
        data = altered(data, {"word/_rels/document.xml.rels": content.encode()})
    elif kind == "missing-content-type":
        data = altered(data, {"[Content_Types].xml": b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>'})
    elif kind == "missing-root-relation":
        data = altered(data, {"_rels/.rels": b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>'})
    else:
        raw = bytearray(data)
        central, end = raw.index(b"PK\x01\x02"), raw.rindex(b"PK\x05\x06")
        if kind == "encrypted":
            struct.pack_into("<H", raw, central + 8, 1)
        elif kind == "zip64":
            struct.pack_into("<I", raw, central + 24, 0xffffffff)
        elif kind == "multivolume":
            struct.pack_into("<H", raw, end + 4, 1)
        elif kind == "compression":
            struct.pack_into("<H", raw, central + 10, 12)
        elif kind == "crc":
            struct.pack_into("<I", raw, central + 16, 0)
        elif kind == "size":
            struct.pack_into("<I", raw, central + 24, 1)
        data = bytes(raw)
    result = await parser.inspect_docx(data)
    assert result["parser_status"] == "parse_failed", (kind, result)
    assert result["preview_text"] == "" and result["parser_reason"] in {
        "invalid_package", "unsafe_package", "package_limit", "unsupported_feature"}


async def test_external_hyperlink_text_is_inert_and_no_connection():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    requests = []
    async def received(reader, writer):
        requests.append(await reader.read(100))
        writer.close()
    server = await asyncio.start_server(received, "127.0.0.1", 0)
    try:
        port = server.sockets[0].getsockname()[1]
        value = document()
        p = value.add_paragraph()
        rel = p.part.relate_to(f"http://127.0.0.1:{port}/must-not-fetch",
                              "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), rel)
        run, text = OxmlElement("w:r"), OxmlElement("w:t")
        text.text = "链接可见文字"
        run.append(text)
        link.append(run)
        p._p.append(link)
        result = await parser.inspect_docx(encoded(value))
        assert result["parser_status"] == "parsed" and "链接可见文字" in result["preview_text"]
        assert not requests
    finally:
        server.close()
        await server.wait_closed()


def probe(monkeypatch, tmp_path, expression):
    # Test-only copy of the actual trusted bootstrap: use identical gates and
    # replace only the post-gate parser call. Production accepts no such mode.
    original = Path(parser._WORKER)
    source = original.read_text(encoding="utf8").replace("__file__", repr(str(original)))
    source = source.replace("result = parse_docx(data)", expression)
    worker = tmp_path / "resource_probe.py"
    worker.write_text(source, encoding="utf8")
    monkeypatch.setattr(parser, "_WORKER", worker)
    return original


@pytest.mark.parametrize("kind", ["cpu", "memory", "stdout", "stderr", "crash", "timeout", "pid"])
async def test_native_resource_child_is_reaped_and_next_parse_healthy(monkeypatch, tmp_path, kind):
    expressions = {"cpu": "exec('while True: pass')", "memory": "result = bytearray(800 * 1024 * 1024)",
                   "stdout": "os.write(1, b'x' * (600 * 1024)); os._exit(0)",
                   "stderr": "os.write(2, b'x' * (32 * 1024)); os._exit(0)",
                   "crash": "os._exit(23)", "timeout": "__import__('time').sleep(60)",
                   "pid": "result = {'pid': os.getpid()}"}
    original = probe(monkeypatch, tmp_path, expressions[kind])
    processes, popen = [], subprocess.Popen
    def launched(*args, **kwargs):
        child = popen(*args, **kwargs)
        processes.append(child)
        return child
    monkeypatch.setattr(parser.subprocess, "Popen", launched)
    if kind == "cpu":
        monkeypatch.setattr(parser, "WALL_SECONDS", 40)
    if kind == "timeout":
        monkeypatch.setattr(parser, "WALL_SECONDS", .3)
    if kind == "pid":
        # Exact PID of actual Windows worker, not the venv redirector PID.
        monkeypatch.setattr(parser, "validate", lambda value: value)
    started = time.monotonic()
    result = await parser.inspect_docx(b"")
    assert len(processes) == 1 and processes[0].poll() is not None
    if kind == "pid":
        assert result == {"pid": processes[0].pid}
    elif kind in {"stdout", "stderr"}:
        assert result["parser_reason"] == "output_limit", result
    elif kind == "timeout":
        assert result["parser_reason"] == "parser_timeout", result
    else:
        assert result["parser_reason"] == "parser_crash", result
    assert time.monotonic() - started < (40 if kind == "cpu" else 14)
    monkeypatch.setattr(parser, "_WORKER", original)
    monkeypatch.setattr(parser, "WALL_SECONDS", 15)
    if kind == "pid":
        from minimax_code.attachment_metadata import validate
        monkeypatch.setattr(parser, "validate", validate)
    assert (await parser.inspect_docx(encoded(document())))["parser_status"] == "parsed"


async def test_native_child_repeated_cancel_reaps_before_return(monkeypatch, tmp_path):
    original = probe(monkeypatch, tmp_path, "__import__('time').sleep(60)")
    processes, popen = [], subprocess.Popen
    def launched(*args, **kwargs):
        child = popen(*args, **kwargs)
        processes.append(child)
        return child
    monkeypatch.setattr(parser.subprocess, "Popen", launched)
    task = asyncio.create_task(parser.inspect_docx(b""))
    for _ in range(100):
        if processes:
            break
        await asyncio.sleep(.01)
    task.cancel()
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await asyncio.wait_for(task, 3)
    assert processes and processes[0].poll() is not None
    monkeypatch.setattr(parser, "_WORKER", original)
    assert (await parser.inspect_docx(encoded(document())))["parser_status"] == "parsed"


def _linux_parser_finished(pid):
    try:
        state = Path(f"/proc/{pid}/stat").read_text().rsplit(")", 1)[1].split()[0]
    except (FileNotFoundError, ProcessLookupError):
        # Linux can remove the process after open(/proc/PID/stat) but before read.
        return True
    return state == "Z"


def test_linux_parser_finished_accepts_only_process_disappearance(monkeypatch):
    def missing(_):
        raise FileNotFoundError()
    def disappeared(_):
        raise ProcessLookupError()
    def denied(_):
        raise PermissionError()
    for read in (missing, disappeared):
        monkeypatch.setattr(Path, "read_text", read)
        assert _linux_parser_finished(1)
    monkeypatch.setattr(Path, "read_text", denied)
    with pytest.raises(PermissionError):
        _linux_parser_finished(1)
    monkeypatch.setattr(Path, "read_text", lambda _: "1 (parser) Z 0")
    assert _linux_parser_finished(1)
    monkeypatch.setattr(Path, "read_text", lambda _: "1 (parser) R 0")
    assert not _linux_parser_finished(1)


def test_native_service_death_terminates_actual_parser(monkeypatch, tmp_path):
    pidfile = tmp_path / "child-pid"
    probe(monkeypatch, tmp_path,
          f"open({str(pidfile)!r}, 'w').write(str(os.getpid())); __import__('time').sleep(60)")
    root = str(Path(__file__).resolve().parents[1])
    paths = sysconfig.get_paths()
    code = (f"import sys; sys.path.extend({[root, paths['purelib'], paths['platlib']]!r}); "
            "import asyncio; from pathlib import Path; from minimax_code import attachment_parser as p; "
            f"p._WORKER=Path({str(parser._WORKER)!r}); asyncio.run(p.inspect_docx(b''))")
    executable = getattr(sys, "_base_executable", sys.executable) if os.name == "nt" else sys.executable
    # This isolated fixture also ignores PYTHONDONTWRITEBYTECODE; keep its imports read-only.
    service = subprocess.Popen([executable, "-I", "-B", "-S", "-c", code],
                               stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                               env=parser._environment(), creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
    try:
        deadline = time.monotonic() + 8
        while not pidfile.exists() and service.poll() is None and time.monotonic() < deadline:
            time.sleep(.02)
        assert pidfile.exists(), service.communicate(timeout=1)
        pid = int(pidfile.read_text())
        if os.name == "nt":
            from minimax_code._windows_process_host import _process_kernel
            kernel = _process_kernel()
            handle = kernel.OpenProcess(0x00100000, False, pid)
            assert handle
            service.kill()
            service.wait(timeout=3)
            try:
                assert kernel.WaitForSingleObject(handle, 3000) == 0
            finally:
                kernel.CloseHandle(handle)
        else:
            service.kill()
            service.wait(timeout=3)
            deadline = time.monotonic() + 3
            while time.monotonic() < deadline:
                if _linux_parser_finished(pid):
                    break
                time.sleep(.02)
            else:
                pytest.fail("Parser survived service death")
    finally:
        if service.poll() is None:
            service.kill()
        service.wait(timeout=3)
        service.stdout.close()
        service.stderr.close()


@pytest.mark.parametrize("kind", ["invalid", "missing_dependency", "limits_unavailable", "timeout"])
async def test_original_retained_with_hash_on_parser_failure(monkeypatch, tmp_path, kind):
    data = encoded(document())
    if kind == "invalid":
        data = b"PK\x03\x04invalid original"
    elif kind == "missing_dependency":
        probe(monkeypatch, tmp_path, "result = parse_docx(data)")
        worker = parser._WORKER
        source = worker.read_text().replace("sys.path.extend(sys.argv[-2:])", "pass")
        worker.write_text(source)
    elif kind == "limits_unavailable":
        probe(monkeypatch, tmp_path, "result = parse_docx(data)")
        worker = parser._WORKER
        worker.write_text(worker.read_text().replace("job = establish_limits()", "raise RuntimeError('gate failed')"))
    else:
        probe(monkeypatch, tmp_path, "__import__('time').sleep(60)")
        monkeypatch.setattr(parser, "WALL_SECONDS", .3)
    db = AsyncDatabase(tmp_path / "data.db")
    await db.connect()
    await db.migrate()
    try:
        await ProjectsDAO(db).create(id="p", name="p")
        service = AttachmentService(db, tmp_path)
        identity = (await service.list("p"))["service_id"]
        async def chunks():
            yield data
        row = await service.upload("p", "original.docx", identity, chunks())
        assert row["parser_status"] in {"parse_failed", "parser_unavailable"}
        assert row["sha256"] == hashlib.sha256(data).hexdigest()
        assert (await service.content("p", row["id"]))[1] == data
        assert (await service.get("p", row["id"]))["preview_text"] == ""
        if kind != "invalid":
            assert row["parser_reason"] == {"timeout": "parser_timeout"}.get(kind, kind)
    finally:
        await db.close()


def test_parent_revalidates_output_caps_and_wire_shape():
    from minimax_code.attachment_metadata import failure, validate
    for change in ({"preview_text": "x" * 65537}, {"preview_text": "x\n" * 2001},
                   {"parser_version": 1}, {"parser_scope": "arbitrary"}, {"parser_reason": "secret"},
                   {"preview_truncated": 1}, {"extra": "forbidden"}):
        with pytest.raises(ValueError):
            validate({**failure("parser_crash"), **change})


async def test_queue_is_single_slot_and_bounded(monkeypatch):
    assert parser._GATE.acquire(timeout=1)
    monkeypatch.setattr(parser, "QUEUE_SECONDS", .1)
    try:
        result = await parser.inspect_docx(b"")
        assert result["parser_reason"] == "queue_timeout"
    finally:
        parser._GATE.release()
