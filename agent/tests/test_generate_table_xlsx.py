"""Ordinary-chat XLSX generation and binding, without model or network calls."""
from __future__ import annotations

from unittest.mock import AsyncMock

import pytest

from minimax_code.agent.core import AgentCore
from minimax_code.agent.tools import ToolRegistry
from minimax_code.agent.tools.base import ToolResult
from minimax_code.agent.tools.generate_table_xlsx import register_table_xlsx_tool
from minimax_code.execution import ExecutionPolicy, execution_scope
from minimax_code.storage.dao.projects import ProjectsDAO
from minimax_code.storage.dao.sessions import SessionsDAO
from minimax_code.storage.db import AsyncDatabase


@pytest.fixture
async def table_env(tmp_path):
    db = AsyncDatabase(tmp_path / "data" / "test.db")
    await db.connect()
    await db.migrate()
    await ProjectsDAO(db).create(id="p", name="Project", root_path=str(tmp_path))
    await ProjectsDAO(db).create(id="other", name="Other", root_path=str(tmp_path))
    sessions = SessionsDAO(db)
    await sessions.create(id="s", project_id="p")
    registry = ToolRegistry()
    register_table_xlsx_tool(registry, db, "s", "p")
    yield db, sessions, registry
    await db.close()


def args():
    return {"filename": "report.xlsx", "request_key": "intent_1", "spec": {
        "sheet_name": "数据", "columns": [{"name": "编号", "kind": "text"},
                                          {"name": "数量", "kind": "number"}],
        "rows": [["00001", 2], ["=literal", 3]]}}


async def test_xlsx_tool_real_generation_retry_and_project_binding(table_env):
    db, sessions, registry = table_env
    with execution_scope(ExecutionPolicy(read_only=True)):
        assert not (await registry.dispatch("generate_table_xlsx", args())).success
    assert not (await registry.dispatch("generate_table_xlsx", {**args(), "project_id": "other"})).success
    first = await registry.dispatch("generate_table_xlsx", args())
    assert first.success and first.output["structure_status"] == "passed"
    assert first.output["business_quality"] == "unverified"
    replay = await registry.dispatch("generate_table_xlsx", args())
    assert replay.success and replay.output["delivery_id"] == first.output["delivery_id"]
    changed = args()
    changed["spec"]["rows"] = [["other", 1]]
    assert not (await registry.dispatch("generate_table_xlsx", changed)).success
    assert (await db.fetchone("SELECT COUNT(*) AS n FROM office_deliveries"))["n"] == 1
    await sessions.update("s", project_id="other")
    assert not (await registry.dispatch("generate_table_xlsx", args())).success


async def test_rebinding_during_generation_cannot_commit(table_env, monkeypatch):
    from minimax_code import attachment_parser

    db, sessions, registry = table_env
    original = attachment_parser.inspect_xlsx

    async def rebind(data):
        parsed = await original(data)
        await sessions.update("s", project_id="other")
        return parsed

    monkeypatch.setattr(attachment_parser, "inspect_xlsx", rebind)
    assert not (await registry.dispatch("generate_table_xlsx", args())).success
    assert (await db.fetchone("SELECT COUNT(*) AS n FROM office_deliveries"))["n"] == 0
    assert (await db.fetchone("SELECT COUNT(*) AS n FROM attachments"))["n"] == 0


@pytest.mark.parametrize(("rules", "has_gater", "consent", "permitted"), [
    ({}, False, False, False), ({}, True, False, False), ({}, True, True, True),
    ({"write_file": "deny", "generate_table_xlsx": "allow"}, True, True, False),
    ({"write_file": "ask", "generate_table_xlsx": "allow"}, True, False, False),
])
async def test_write_permissions(table_env, rules, has_gater, consent, permitted):
    _, _, registry = table_env
    tool = registry.get("generate_table_xlsx")
    tool.run = AsyncMock(return_value=ToolResult.ok())

    class Store:
        def lookup(self, name):
            return {"action": rules[name]} if name in rules else None

    gater = type("Gater", (), {"request_consent": AsyncMock(return_value=consent)})() if has_gater else None
    core = AgentCore(llm=object(), registry=registry, permission_store=Store(), permission_gater=gater)
    result = await core._dispatch_tool({"name": "generate_table_xlsx", "arguments": args()})
    assert result.success is permitted
    assert tool.run.await_count == (1 if permitted else 0)


async def test_summaries_and_chart_are_real_xlsx_values_and_attachment(table_env):
    from io import BytesIO
    from zipfile import ZIP_DEFLATED, ZipFile

    import openpyxl

    from minimax_code.attachment_docx import Reject
    from minimax_code.attachments import AttachmentService
    from minimax_code.office_analysis import AnalysisSpec, Chart, Summary, materialize, verify_chart
    from minimax_code.office_xlsx import XlsxSpec, verify

    db, _, registry = table_env
    request = args()
    request["spec"]["rows"] = [["甲", 1], ["乙", None], ["丙", 2]]
    request["summaries"] = [{"column": "数量", "operation": "sum"},
                            {"column": "编号", "operation": "count"},
                            {"column": "数量", "operation": "average", "decimals": 2}]
    request["chart"] = {"kind": "bar", "category_column": "编号", "value_column": "数量", "title": "数量图"}
    result = await registry.dispatch("generate_table_xlsx", request)
    assert result.success, result.error
    assert result.output["validation_scope"] == "xlsx_chart_values_v1"
    assert result.output["chart_reference_check"] == "passed"
    service = AttachmentService(db, db.path.parent)
    row, data = await service.content("p", result.output["attachment_id"])
    assert row["parser_status"] in {"parsed", "truncated"}
    sheet = openpyxl.load_workbook(BytesIO(data)).active
    assert [sheet.cell(row, 4).value for row in (5, 6, 7)] == [3, 3, 1.5]
    assert len(sheet._charts) == 1
    original = XlsxSpec.model_validate(request["spec"])
    expanded = materialize(original, [Summary.model_validate(item) for item in request["summaries"]])
    analysis = AnalysisSpec(spec=expanded, chart=Chart.model_validate(request["chart"]), source_rows=3)
    assert verify_chart(data, analysis)["cell_check"] == "passed"
    tampered = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(tampered, "w", ZIP_DEFLATED) as target:
        for name in source.namelist():
            value = source.read(name)
            if name == "xl/charts/chart1.xml":
                assert b"$B$2:$B$4" in value
                value = value.replace(b"$B$2:$B$4", b"$B$2:$B$3")
            target.writestr(name, value)
    with pytest.raises(Reject):
        verify_chart(tampered.getvalue(), analysis)
    with pytest.raises(Reject):
        verify(data, expanded)  # The historical plain-cell validator never accepts chart packages.
    assert (await db.fetchone("SELECT COUNT(*) AS n FROM office_deliveries"))["n"] == 0


def test_summary_null_semantics_and_rounding():
    from minimax_code.office_analysis import Summary, materialize
    from minimax_code.office_xlsx import XlsxSpec

    spec = XlsxSpec.model_validate({"sheet_name": "表", "columns": [{"name": "数", "kind": "number"}],
                                    "rows": [[None], [1], [2]]})
    output = materialize(spec, [Summary(column="数", operation="count"),
                                Summary(column="数", operation="sum"),
                                Summary(column="数", operation="average", decimals=0)])
    assert [row[-1] for row in output.rows[-3:]] == [2, 3, 2]
    empty = XlsxSpec.model_validate({"sheet_name": "表", "columns": [{"name": "数", "kind": "number"}], "rows": [[None]]})
    output = materialize(empty, [Summary(column="数", operation="sum"), Summary(column="数", operation="average")])
    assert [row[-1] for row in output.rows[-2:]] == [0, None]


async def test_summary_only_uses_plain_receipt_and_replays(table_env):
    from io import BytesIO

    import openpyxl

    from minimax_code.office_deliveries import OfficeService

    db, _, registry = table_env
    request = args()
    request["summaries"] = [{"column": "数量", "operation": "sum"}]
    first = await registry.dispatch("generate_table_xlsx", request)
    assert first.success and first.output["validation_scope"] == "xlsx_plain_cells_v1"
    again = await registry.dispatch("generate_table_xlsx", request)
    assert again.success and again.output["delivery_id"] == first.output["delivery_id"]
    service = OfficeService(db, db.path.parent)
    _, content = await service.originals.content("p", first.output["attachment_id"])
    sheet = openpyxl.load_workbook(BytesIO(content)).active
    assert sheet["C4"].value == "sum(数量)" and sheet["D4"].value == 5
    request["summaries"] = [{"column": "数量", "operation": "average"}]
    assert not (await registry.dispatch("generate_table_xlsx", request)).success
